import datetime
import os
import sys
import traceback
import tempfile
from os import listdir
from os.path import isfile, join
# SUBPROCESS (EXE)
import subprocess

# CONFIGURAZIONE CIFRATA
from config_sicuro import carica_config

# EMAIL
import email
import smtplib
import ssl
from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

# INTERNET (TELEGRAM)
import requests
import telepot

#import cookie

# WHATSAPP (API Cloud di Meta)
from whatsapp_meta import WhatsApp, WhatsAppError

# PDF
from PyPDF2 import PdfFileWriter, PdfFileReader
from fpdf import FPDF
import fpdf.fpdf
# Stampa senza Ghostscript (PyMuPDF + stampante di Windows)
from stampa_pdf import stampa_pdf

# DOCX
from docx import Document
from docx.shared import RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, Cm

# EXCEL
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

#BARCODE
import barcode
from barcode.writer import ImageWriter
import fitz

import time

def process(filesprec):
    # Legge la configurazione cifrata (config.dat, si modifica con ConfigManager.py)
    config = carica_config()

    # Cartella con i file di testo e i .gui
    folder = config["cartella_lavoro"]

    # Legge i file in questa cartella
    onlyfiles = [f for f in listdir(folder) if isfile(join(folder, f))]

    if filesprec != onlyfiles:
        start()
    return onlyfiles

def main(name, config):
    # Sostituisce al file .txt la terminazione .gui per vedere se esiste
    name_gui = name.replace(".txt", ".gui")
    path_gui = config["cartella_lavoro"] + "\\" + name_gui
    if name_gui[len(name_gui) - 3:len(name_gui)] != "gui":
        return
    if not os.path.isfile(path_gui):
        return
    # Apre il file .gui
    f = open(path_gui, "r")
    guida = f.readlines()
    f.close()
    if guida[1].replace("\n", "") == "exe":
        exe(guida, config)
        # Rimuove i file dopo l'uso
        os.remove(path_gui)
    elif guida[1].replace("\n", "") == "barcode":
        barcode_add(guida, config)
        os.remove(path_gui)
    elif guida[1].replace("\n", "") == "immagini":
        image_add(guida, config)
        os.remove(path_gui)

    if name[-3:] == "gui":
        return

    # Legge la seconda linea del .gui per decidere cosa fare
    if guida[1].replace("\n", "") == "pdf" or guida[1].replace("\n", "") == "file" \
            or guida[1].replace("\n", "") == "stampa":
        pdf(guida, config)
        if guida[5].replace("\n", "") == "email":
            email_send(guida, config)
        if guida[5].replace("\n", "") == "telegram":
            telegram(guida, config)
        if guida[5].replace("\n", "") == "whatsapp":
            whatsapp(guida, config)
    if guida[1].replace("\n", "") == "excel":
        excel(guida, config)
        if guida[5].replace("\n", "") == "email":
            email_send(guida, config)
        if guida[5].replace("\n", "") == "telegram":
            telegram(guida, config)
        if guida[5].replace("\n", "") == "whatsapp":
            whatsapp(guida, config)
    if guida[1].replace("\n", "") == "docx":
        docx(guida, config)
        if guida[5].replace("\n", "") == "email":
            email_send(guida, config)
        if guida[5].replace("\n", "") == "telegram":
            telegram(guida, config)
        if guida[5].replace("\n", "") == "whatsapp":
            whatsapp(guida, config)

    # Rimuove i file dopo l'uso
    os.remove(path_gui)
    os.remove(config["cartella_lavoro"] + "\\" + guida[0].replace("\n", ""))


def barcode_add(guida, config):
    numero_barcodes = int(guida[2].replace("\n", ""))
    for i in range(0, numero_barcodes):
        i *= 6
        codice = guida[3 + i].replace("\n", "")
        cab = barcode.get_barcode_class(guida[4 + i].replace("\n", ""))
        codicino = cab(codice, writer=ImageWriter())
        percorso = config["cartella_lavoro"] + "\\" + "barcode" + str(codice)
        codicino.save(percorso)

        # la posizione è data da due vertici opposti partendo da quello alto a sx
        x0, y0 = guida[5 + i].replace("\n", "").split(",")
        x1, y1 = guida[6 + i].replace("\n", "").split(",")
        image_rectangle = fitz.Rect(int(x0), int(y0), int(x1), int(y1))

        # retrieve the page of the PDF
        doc = fitz.open(guida[7 + i].replace("\n", ""))
        page = doc[int(guida[8 + i].replace("\n", ""))]

        # add the image
        page.insert_image(image_rectangle, filename=percorso + ".png")
        doc.save(config["cartella_lavoro"] + "\\" + guida[7 + i].replace("\n", "").split("\\")[-1], incremental=True, encryption=0)
        doc.close()
        #L'immagine del barcode non serve più
        os.remove(percorso + ".png")

def image_add(guida,config):
    numero_image = int(guida[2].replace("\n", ""))
    for i in range(0, numero_image):
        i *= 5
        # la posizione è data da due vertici opposti partendo da quello alto a sx
        percorso = guida[3 + i].replace("\n", "")
        x0, y0 = guida[4 + i].replace("\n", "").split(",")
        x1, y1 = guida[5 + i].replace("\n", "").split(",")
        image_rectangle = fitz.Rect(int(x0), int(y0), int(x1), int(y1))

        # retrieve the page of the PDF
        doc = fitz.open(guida[6 + i].replace("\n", ""))
        page = doc[int(guida[7 + i].replace("\n", ""))]

        # add the image
        page.insert_image(image_rectangle, filename= percorso)
        doc.save(config["cartella_lavoro"] + "\\" + guida[6 + i].replace("\n", "").split("\\")[-1], incremental=True, encryption=0)
        doc.close()



def exe(guida, config) :

    path = config["cartella_lavoro"] + "\\" + guida[0].replace("\n", "")
    if not os.path.isfile(path):
        path = guida[0].replace("\n", "")
    if not os.path.isfile(path):
        return
    try:
        #Tra virgolette, così funzionano anche i percorsi con spazi
        os.system('"' + path + '"')
    except:
        subprocess.call(path)

def telegram(guida, config):
    #Legge il token del bot
    #Per creare un bot si usa TheBotFather su telegram
    bot_token = config["telegram_token"]
    #Quando si manda un messaggio al bot da queso link se ne ottengono i dati
    # https://api.telegram.org/bot<bot_token>/getUpdates
    i = 0
    body = ""
    document = ""
    #argomenti della libreria telepot
    TOKEN = bot_token
    bot = telepot.Bot(TOKEN)

    #legge le ultime linee del file dopo "fine"
    #la prima linea è l'url del documento
    #la seconda e seguenti sono il testo del messaggio

    for x in guida:
        if i == 1 or i == 2:
            if i == 1:
                document = x.replace("\n", "")
            i += 1
        if i >= 3:
            body = body + " " + x
        if x[0:4] == "fine":
            i += 1
    bot_message = body
    bot_chatID = guida[6].replace("\n", "").replace("telegram ", "")
    doc_path = config["cartella_output"] + "\\" + document

    #manda il messaggio
    bot.sendMessage(bot_chatID, bot_message)
    #se esiste manda il documento
    if document != "":
        bot.sendDocument(chat_id=bot_chatID, document=open(doc_path, 'rb'))

def whatsapp(guida, config):
    #Usa le API Cloud di Meta: token e Phone number ID si impostano con ConfigManager.py
    #La settima linea del .gui è "whatsapp <numero> [<template> [<lingua>]]"
    #il numero va con il prefisso internazionale (es. 393331234567 o +39 333 1234567)
    #legge le ultime linee del file dopo "fine"
    #la prima linea è il nome del documento nella cartella di output (vuota = nessun documento)
    #la seconda e seguenti sono il testo del messaggio
    #Senza template il messaggio arriva solo se il destinatario ha scritto al numero
    #aziendale nelle ultime 24 ore (regola di Meta).
    #Con il template le linee di testo diventano i parametri {{1}}, {{2}}, ... del template
    #e il documento diventa l'intestazione (il template deve avere un'intestazione Documento)
    parti = guida[6].replace("\n", "").split()[1:]
    #il numero può essere scritto con spazi: prende tutte le parti fatte solo di cifre e +-./()
    numero = ""
    while len(parti) > 0 and parti[0].strip("+-./()0123456789") == "":
        numero += parti.pop(0)
    if numero == "":
        print("WhatsApp: numero del destinatario mancante nella linea 7 del .gui")
        return
    template = parti[0] if len(parti) > 0 else ""
    lingua = parti[1] if len(parti) > 1 else config["whatsapp_lingua"]

    #cerca la linea "fine" e prende le linee successive
    fine = len(guida)
    for n in range(linea_inizio_format_pdf, len(guida)):
        if guida[n][0:4] == "fine":
            fine = n
            break
    coda = [x.replace("\n", "") for x in guida[fine + 1:]]
    document = coda[0].strip() if len(coda) > 0 else ""
    righe = coda[1:]
    doc_path = config["cartella_output"] + "\\" + document if document != "" else None

    try:
        messenger = WhatsApp(config["whatsapp_token"], config["whatsapp_phone_number_id"],
                             config["whatsapp_versione_api"])
        if template != "":
            parametri = [r for r in righe if r.strip() != ""]
            messenger.invia_template(numero, template, lingua, parametri, doc_path)
        else:
            testo = "\n".join(righe).strip()
            #manda il messaggio
            if testo != "":
                messenger.invia_testo(numero, testo)
            #se esiste manda il documento
            if doc_path is not None:
                messenger.invia_documento(numero, doc_path)
    except (WhatsAppError, requests.RequestException, OSError) as e:
        print("Errore invio WhatsApp a " + numero + ": " + str(e))

def pdf(guida, config):

    #Apre il documento di testo
    f = open(config["cartella_lavoro"] + "\\" + guida[0].replace("\n", ""), 'r')
    content = f.readlines()
    f.close()

    #Immagazzina la rotazione(eventuale) da applicare alla pagina
    rotazione = int(guida[4])

    #Dichiara il documento
    #Se il template è orizzontale anche il testo va su una pagina orizzontale,
    #altrimenti le prime righe finiscono fuori dal foglio
    orientamento = "P"
    nome_templ = guida[3].replace("\n", "")
    if nome_templ != "":
        with fitz.open(config["cartella_template"] + "\\" + nome_templ) as doc_templ:
            if doc_templ[0].mediabox.width > doc_templ[0].mediabox.height:
                orientamento = "L"
    pdf = FPDF(orientation=orientamento)

    #Imposta font, dimensione e colore. Il font viene aggiunto al pdf solo la prima volta che serve
    #(ogni font aggiunto finisce dentro il file). Lo cerca nella cartella Fonts del programma,
    #se manca usa quello di Windows (Helvetica in Windows non c'è: al suo posto usa Arial)
    def usa_font(nome_font, dimensione, colore="n"):
        if nome_font != "" and nome_font.lower() not in pdf.fonts:
            file_font, file_windows = FONT_PDF[nome_font]
            percorso_font = os.path.join(CARTELLA_FONT, file_font)
            if not os.path.isfile(percorso_font):
                percorso_font = os.path.join(CARTELLA_FONT_WINDOWS, file_windows)
            pdf.add_font(nome_font, '', percorso_font, uni=True)
        pdf.set_font(family=nome_font, size=dimensione)
        pdf.set_text_color(*COLORI_TESTO.get(colore, (0, 0, 0)))

    #Aggiunge la prima pagina al pdf
    pdf.alias_nb_pages()
    pdf.add_page()
    #Font predefinito per le righe senza formato (stesso del codice iniziale "tnn10")
    usa_font("Times", 10)

    i = linea_inizio_format_pdf
    index = 0
    testopasprec = ""
    testopass = ""
    codice = "tnn10"
    #For praticamente identico a quello della funzione word
    #vedi nota lì
    for x in content:
        arrcodici = []
        riga = x
        car1prec = 0
        #Il testo prima del formato va azzerato a ogni riga, altrimenti si ripete nelle righe dopo
        testopasprec = ""
        if guida[i][0:4] == "fine":
            i = linea_inizio_format_pdf
            index = 0

        codicecompleto = guida[i].replace("\n", "")[4:]
        #Il codice ha 5 caratteri (es. tsn12) o 6 con il colore (es. tsn12r)
        lunghezza_codice = 6 if codicecompleto[5:6] in COLORI_TESTO else 5
        if len(codicecompleto.rstrip()) > lunghezza_codice:
            #Più formati sulla riga: "<codice> <da> <a>" separati da un carattere
            passo = lunghezza_codice + 7
            times = int((len(codicecompleto) + 1) / passo)
            for h in range(0, times):
                primochar = passo*h
                ultimochar = passo*h + lunghezza_codice + 6
                arrcodici.append(codicecompleto[primochar: ultimochar])
        else:
            arrcodici.append(codicecompleto)

        index += 1
        testo_tot = ""
        def controllacodice(codice):
                #Font (es. Times, TimesBold, TimesBoldIt, TimesIt), dimensione e colore
                font = ""
                if codice[0:1].lower() in FONT_CODICI:
                    font = FONT_CODICI[codice[0:1].lower()][0]
                    if codice[1:2].lower() == "s":
                        font += "Bold"
                    if codice[2:3].lower() == "s":
                        font += "It"
                size = codice[3:5]
                colore = codice[5:6] if codice[5:6] in COLORI_TESTO else "n"
                array = [font, size, colore]
                return array

        if arrcodici[0] != "" and arrcodici[0] != " ":
            for codici in arrcodici:
                codice2 = codici[0:lunghezza_codice]
                if codici[lunghezza_codice + 1:] != "":
                    car1 = int(codici[lunghezza_codice + 1:lunghezza_codice + 4])
                    car2 = int(codici[lunghezza_codice + 4:])
                    testopass = riga[car1: car2]
                    testo_tot += testopass
                    if codici == arrcodici[0]:
                        testopasprec = riga[0: car1]
                        testo_tot += testopasprec
                    else:
                        testopasprec = riga[car1prec: car1]
                        testo_tot += testopasprec
                    car1prec = car2
                else:
                    testopass = riga
                    testo_tot += testopass
                if testopasprec != "":
                    testoarr = controllacodice(codice)
                    codice = guida[i][4:4 + lunghezza_codice]
                    testo = testopasprec.replace("\n", "")
                    usa_font(testoarr[0], int(testoarr[1]), testoarr[2])
                    #Larghezza della cella = larghezza reale del testo, così i pezzi non si sovrappongono
                    pdf.cell(pdf.get_string_width(testo), 6, testo, border=0, ln=0)
                testoarr2 = controllacodice(codice2)
                testo = testopass.replace("\n", "")
                usa_font(testoarr2[0], int(testoarr2[1]), testoarr2[2])
                if codici == arrcodici[len(arrcodici) - 1]:
                    #Il testo dopo l'ultimo pezzo formattato prende lo stesso formato
                    if codici[lunghezza_codice + 1:] != "":
                        testo += riga[car1prec:].replace("\n", "")
                    pdf.cell(0, 6, testo, border=0, ln=1)
                else:
                    pdf.cell(pdf.get_string_width(testo), 6, testo, border=0, ln=0)
        else:
            testo = x.replace("\n", "")
            pdf.cell(0, 6, testo, border=0, ln=1)
        i += 1

    #Salva il pdf temporaneo
    pdf.output(config["cartella_lavoro"] + "\\" + "systemgen.pdf")
    pdf.close()

    #Prepara i pdf overlay template e result
    nome_templ = guida[3].replace("\n", "")
    overlay_pdf_file_name = config["cartella_lavoro"] + "\\" + "systemgen.pdf"

    with open(overlay_pdf_file_name, 'rb') as Pdf_overlay:
        if nome_templ != "":
            pdf_template_file_name = config["cartella_template"] + "\\" + nome_templ
            result_pdf_file_name = config["cartella_output"] + "\\" + guida[2].replace("\n", "")

            # Open your overlay PDF that was created earlier

            overlay_pdf = PdfFileReader(Pdf_overlay)
            numpag = overlay_pdf.getNumPages()

            # Merge the overlay page onto the template page
            output_pdf = PdfFileWriter()
            #Il template resta aperto finché il risultato non è scritto (PyPDF2 lo legge in quel momento)
            file_template = []
            for intero in range(0, numpag):
                file_template.append(open(pdf_template_file_name, 'rb'))
                pdf_template = PdfFileReader(file_template[-1])
                template_page = pdf_template.getPage(0)
                template_page.rotateClockwise(rotazione)
                overlay_page = overlay_pdf.getPage(intero)
                template_page.mergePage(overlay_page)

                # Write the result to a new PDF file
                output_pdf.addPage(template_page)

            # Salva tutto nel result
            with open(result_pdf_file_name, 'wb') as resultPdf:
                output_pdf.write(resultPdf)
            for f in file_template:
                f.close()

        else:
            result_pdf_file_name = config["cartella_output"] + "\\" + guida[2].replace("\n", "")
            output_pdf = PdfFileWriter()
            # Open your overlay PDF that was created earlier

            overlay_pdf = PdfFileReader(Pdf_overlay)
            numpag = overlay_pdf.getNumPages()

            for intero in range(0, numpag):
                over_page = overlay_pdf.getPage(intero)
                over_page.rotateClockwise(rotazione)

                # Write the result to a new PDF file
                output_pdf.addPage(over_page)

            # Salva tutto nel result
            with open(result_pdf_file_name, 'wb') as resultPdf:
                output_pdf.write(resultPdf)

    #Il risultato è già chiuso (scritto del tutto) prima di aprirlo o stamparlo
    if guida[1].replace("\n", "") == "pdf":
        os.startfile(result_pdf_file_name)
    if guida[1].replace("\n", "") == "stampa":
        stampa_pdf(result_pdf_file_name)
    #Prova a eliminare il file temporaneo
    try:
        os.remove(config["cartella_lavoro"] + "\\" + "systemgen.pdf")
    except:
        print("Can't remove systemgen.pdf")

def docx(guida, config):
    #Apre il file di testo
    f = open(config["cartella_lavoro"] + "\\" + guida[0].replace("\n", ""), 'r')
    content = f.readlines()
    f.close()

    #dichiara il documento
    document = Document()
    sections = document.sections
    for section in sections:
        section.top_margin = Cm(1.4)
        section.bottom_margin = Cm(1.4)
        section.left_margin = Cm(1.6)
        section.right_margin = Cm(1.6)
    par = document.add_paragraph()
    #dichiara lo stile
    style = document.styles.add_style("Custom", WD_STYLE_TYPE.CHARACTER)
    i = linea_inizio_format_pdf
    index = 0
    testo_tot = ""
    testopasprec = ""
    #Formato delle righe senza codice prima del primo formato
    lastfont = "Times New Roman"
    lastsize = Pt(10)
    lastbold = False
    lastit = False
    lastcolor = "n"
    #for presente anche nella funzione pdf
    #per ogni linea del file di testo va a leggere la corrispondente linea del file guida
    #se è presente applica una formattazione
    #aggiunge linea per linea il testo al documento Word
    for x in content:
        riga = x
        car1prec = 0
        arrcodici = []
        #Il testo prima del formato va azzerato a ogni riga, altrimenti si ripete nelle righe dopo
        testopasprec = ""
        #Finiti i formati riparte dal primo, come nel pdf
        if guida[i][0:4] == "fine":
            i = linea_inizio_format_pdf
            index = 0
        if guida[i][0:4] == "fine":
            #Il .gui non ha formati: la riga mantiene l'ultimo formato
            run = par.add_run(riga)
            run.font.name = lastfont
            run.font.size = lastsize
            run.font.bold = lastbold
            run.font.italic = lastit
            run.font.color.rgb = RGBColor(*COLORI_TESTO.get(lastcolor, (0, 0, 0)))
        else:
            codicecompleto = guida[i].replace("\n", "")[4:]
            if guida[i][11:13] != "":
                lung = len(guida[i].replace("\n", "")) - 3
                times = int(lung / 13)
                for h in range(0, times):
                    primochar = 13 * h
                    ultimochar = 13 * h + 12
                    arrcodici.append(codicecompleto[primochar: ultimochar])
            else:
                arrcodici.append(codicecompleto)
            linea = guida[i][0:3]
            codice = guida[i][4:10]
            index += 1

            def controllacodice(codice):
                #Nome del font come lo conosce Word (es. Times New Roman), dimensione, grassetto,
                #corsivo e colore
                font = None
                if codice[0:1].lower() in FONT_CODICI:
                    font = FONT_CODICI[codice[0:1].lower()][1]
                bold = codice[1:2].lower() == "s"
                italic = codice[2:3].lower() == "s"
                size = codice[3:5]
                colore = codice[5:6]
                array = [font, size, bold, italic, colore]
                return array

            if codice != "":
                    for codici in arrcodici:
                        codice2 = codici[0:6]
                        if codici[7:] != "":
                            car1 = int(codici[7:10])
                            car2 = int(codici[10:])
                            testopass = riga[car1: car2]
                            testo_tot += testopass
                            if codici == arrcodici[0]:
                                testopasprec = riga[0: car1]
                                testo_tot += testopasprec
                            else:
                                testopasprec = riga[car1prec: car1]
                                testo_tot += testopasprec
                            car1prec = car2
                        else:
                            testopass = riga
                            testo_tot += testopass
                        if testopasprec != "":
                            testoarr = controllacodice(codice)
                            codice = guida[i][4:10]
                            testo = testopasprec
                            run = par.add_run(testo)
                            font = run.font
                            font.name = testoarr[0]
                            font.size = Pt(int(testoarr[1]))
                            font.bold = testoarr[2]
                            font.italic = testoarr[3]
                            font.color.rgb = RGBColor(*COLORI_TESTO.get(testoarr[4], (0, 0, 0)))

                        testoarr2 = controllacodice(codice2)
                        testo = testopass
                        #Il testo dopo l'ultimo pezzo formattato prende lo stesso formato
                        if codici == arrcodici[len(arrcodici) - 1] and codici[7:] != "":
                            testo += riga[car1prec:]
                        run = par.add_run(testo)
                        font = run.font
                        font.name = testoarr2[0]
                        font.size = Pt(int(testoarr2[1]))
                        font.bold = testoarr2[2]
                        font.italic = testoarr2[3]
                        lastfont = testoarr2[0]
                        lastsize = Pt(int(testoarr2[1]))
                        lastbold = testoarr2[2]
                        lastit = testoarr2[3]
                        lastcolor = testoarr2[4]
                        font.color.rgb = RGBColor(*COLORI_TESTO.get(lastcolor, (0, 0, 0)))
            else:
                testo = riga
                run = par.add_run(testo)
                font = run.font
                font.name = lastfont
                font.size = lastsize
                font.bold = lastbold
                font.italic = lastit
                font.color.rgb = RGBColor(*COLORI_TESTO.get(lastcolor, (0, 0, 0)))
            i += 1


    #Salva il documento
    document.save(config["cartella_output"] + "\\" + guida[2].replace("\n", ""))

def excel(guida, config):
    #Salva in content il contenuto del file di testo
    f = open(config["cartella_lavoro"] + "\\" + guida[0].replace("\n", ""), 'r')
    content = f.readlines()
    f.close()

    #Dichiara l'array delle colonne
    colonne = []

    primo_char = 0

    # Create new workbook
    wb = Workbook()
    sheet1 = wb.create_sheet('sheet1', 0)
    #Toglie il foglio vuoto "Sheet" che Workbook crea da solo
    wb.remove(wb['Sheet'])
    active = wb['sheet1']
    ws = wb.active
    foglio = []


    #legge linea per il linea il file e mette le righe in un array
    for line in content:
        riga = []
        num_col = 1
        i = 0
        primo_char = 0
        while i <= len(line):
            char = line[i: i + 1]
            if char == ";":
                res = isinstance(line[primo_char: i], str)
                if res == True:
                    riga.append(line[primo_char: i])
                else:
                    riga.append(float(line[primo_char: i]))
                primo_char = i + 1
                num_col += 1
            i += 1
        colonne.append(riga)

    #Mette le righe nel foglio
    #(Prima ogni riga passava da un data frame di pandas e se ne prendeva una ogni 3 righe prodotte:
    #una riga vuota ne produceva 2 e faceva slittare tutte le successive)
    for x in range(len(colonne)):
        foglio.append(colonne[x])


    #Formati Excel: nel file xlsx vanno scritti in notazione inglese (Excel li mostra poi
    #con le impostazioni italiane del PC)
    formati_excel = {
        0: '', 1: '0', 2: '0.00', 3: '#,##0', 4: '#,##0.00',
        5: '#,##0 "€";-#,##0 "€"', 6: '#,##0 "€";[Red]-#,##0 "€"', 7: '#,##0 "€";[Red]-#,##0 "€"',
        8: '#,##0.00 "€";[Red]-#,##0.00 "€"', 9: '0%', 10: '0.00%', 11: '0.00E+00',
        12: 'dd/mm/yyyy', 13: 'dd-mmm-yyyy', 14: 'dd-mmm', 15: 'mmm yyyy', 16: 'h:mm AM/PM',
        17: 'h:mm:ss AM/PM', 18: 'hh:mm', 19: 'hh:mm:ss', 20: 'dd/mm/yyyy hh:mm', 21: '##0.0E+0',
        22: 'mm:ss', 23: '@',
    }
    codici_num = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 21]
    codici_data = [12, 13, 14, 15, 20]
    codici_ora = [16, 17, 18, 19, 22]

    def valore_excel(testo, codice):
        #Trasforma il testo nel valore giusto per il formato; se non ci riesce lo lascia come testo
        testo = testo.strip()
        try:
            if codice in codici_num:
                #Numeri all'italiana: 1.234,56 -> 1234.56 (i punti delle migliaia si tolgono)
                numero = testo.replace("€", "").replace("%", "").strip()
                if "," in numero:
                    numero = numero.replace(".", "").replace(",", ".")
                valore = float(numero)
                if codice in [9, 10] and "%" in testo:
                    valore = valore / 100
                return valore
            if codice in codici_data:
                for formato_data in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y"):
                    try:
                        return datetime.datetime.strptime(testo, formato_data)
                    except ValueError:
                        pass
            if codice in codici_ora:
                for formato_ora in ("%H:%M:%S", "%H:%M"):
                    try:
                        return datetime.datetime.strptime(testo, formato_ora).time()
                    except ValueError:
                        pass
        except ValueError:
            pass
        return testo

    #Applica a ogni colonna il formato della sua linea del .gui (dalla linea "fine" in poi: nessun formato)
    #La prima linea dei formati è per la prima colonna (prima andava a una colonna indice
    #che poi veniva cancellata, e tutti i formati risultavano spostati di una colonna)
    i = linea_inizio_format_excel
    arrivato_a_fine = False
    num_colonne = max(len(r) for r in foglio) if len(foglio) > 0 else 0
    for col in range(0, num_colonne):
        linea = guida[i].strip() if i < len(guida) else "fine"
        if linea[0:4] == "fine":
            arrivato_a_fine = True
        try:
            codice = 0 if arrivato_a_fine else int(linea)
        except ValueError:
            codice = 0
        formato = formati_excel.get(codice, '')

        for f in range(0, len(foglio)):
            valore = foglio[f][col] if col < len(foglio[f]) else ""
            if valore is None:
                valore = ""
            #La prima riga è l'intestazione e resta testo
            if f != 0 and isinstance(valore, str):
                valore = valore_excel(valore, codice)
            cella = ws.cell(row=f + 1, column=col + 1)
            cella.value = valore
            if formato != "":
                cella.number_format = formato
        if not arrivato_a_fine:
            i += 1

    #Aggiusta la dimensione delle colonne
    column_widths = []
    for row in sheet1.iter_rows():
        for n, cell in enumerate(row):
            larghezza = len(str(cell.value)) if cell.value is not None else 0
            if n < len(column_widths):
                column_widths[n] = max(column_widths[n], larghezza)
            else:
                column_widths.append(larghezza)
    for n, column_width in enumerate(column_widths):
        sheet1.column_dimensions[get_column_letter(n + 1)].width = column_width + 2

    # Salva il file Excel
    wb.save(config["cartella_output"] + "\\" + guida[2].replace("\n", ""))

def email_send(guida, config):
    i = 0
    body = ""
    subject = ""
    receiver_email = ""
    filename2 = ""
    #legge le linee dopo fine
    #la prima è il ricevente
    #la seconda è l'oggetto
    #la terza è un eventuale secondo allegato
    #Quarta e seguenti sono per il messaggio
    for x in guida:
        if i == 1 or i == 2:
            if i == 1:
                receiver_email = x.replace("\n", "")
            i += 1
        if i == 3 or i == 4:
            if i == 3:
                subject = x.replace("\n", "")
            i += 1
        if i == 5 or i == 6:
            if i == 5:
                filename2 = x.replace("\n", "")
            i += 1
        if i >= 7:
            body = body + " " + x
            i += 1
        if x[0:4] == "fine":
            i += 1

    port = int(config["email_smtp_porta"])  # For SSL

    #passord della mail
    password = config["email_password"]

    # Create a secure SSL context
    context = ssl.create_default_context()

    #L'email che spedisce
    sender_email = config["email_mittente"]

    #L'allegato generato dalla funzione pdf
    filename = config["cartella_output"] + "\\" + guida[2].replace("\n", "")

    if receiver_email == "":
        return

    #Forma il messaggio
    message = MIMEMultipart()
    message["From"] = sender_email
    message["To"] = receiver_email
    message["Subject"] = subject
    message["Bcc"] = receiver_email  # Recommended for mass emails

    # Add body to email
    message.attach(MIMEText(body, "plain"))

    # Open PDF file in binary mode
    with open(filename, "rb") as attachment:
        # Add file as application/octet-stream
        # Email client can usually download this automatically as attachment
        part = MIMEBase("application", "octet-stream")
        part.set_payload(attachment.read())
    # Encode file in ASCII characters to send by email
    encoders.encode_base64(part)
    # Add header as key/value pair to attachment part
    part.add_header(
        "Content-Disposition",
        #Solo il nome del file, non tutto il percorso
        "attachment", filename=os.path.basename(filename),
    )

    # Add attachment to message and convert message to string
    message.attach(part)

    if filename2 != "":
        with open(filename2, "rb") as attachment:
            # Add file as application/octet-stream
            # Email client can usually download this automatically as attachment
            part2 = MIMEBase("application", "octet-stream")
            part2.set_payload(attachment.read())
        # Encode file in ASCII characters to send by email
        encoders.encode_base64(part2)
        # Add header as key/value pair to attachment part
        part2.add_header(
            "Content-Disposition",
            "attachment", filename=os.path.basename(filename2),
        )

        # Add attachment to message and convert message to string
        message.attach(part2)

    #il messaggio che sarà inviato
    text = message.as_string()

    #invia la mail
    with smtplib.SMTP_SSL(config["email_smtp_server"], port, context=context) as server:
        server.login(sender_email, password)
        server.sendmail(sender_email, receiver_email, text)

def start():
    # Legge la configurazione cifrata (config.dat, si modifica con ConfigManager.py)
    config = carica_config()

    # Cartella con i file di testo e i .gui
    folder = config["cartella_lavoro"]

    # Legge i file in questa cartella
    onlyfiles = [f for f in listdir(folder) if isfile(join(folder, f))]

    # Per ogni file passa il suo nome alla main
    for name in onlyfiles:
        #Un file sbagliato non deve bloccare gli altri: scrive l'errore e passa al successivo
        try:
            main(name, config)
        except Exception:
            print("Errore con il file " + name + ":", file=sys.stderr)
            traceback.print_exc()

# file .gui global
linea_inizio_format_pdf = 7
linea_inizio_format_excel = 7


# Font del PDF: cartella Fonts accanto al programma (o all'eseguibile creato con PyInstaller)
if getattr(sys, "frozen", False):
    CARTELLA_FONT = os.path.join(os.path.dirname(sys.executable), "Fonts")
else:
    CARTELLA_FONT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Fonts")
CARTELLA_FONT_WINDOWS = os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts")
# Lettera del font nei codici del .gui -> (nome nel PDF, nome in Word)
FONT_CODICI = {
    "t": ("Times", "Times New Roman"),
    "h": ("Helvetica", "Arial"),
    "c": ("Courier", "Courier New"),
    "k": ("Calibri", "Calibri"),
    "v": ("Verdana", "Verdana"),
    "g": ("Georgia", "Georgia"),
    "u": ("SegoeUI", "Segoe UI"),
}
# Nome usato nel PDF -> (file nella cartella Fonts, file di Windows se manca)
FONT_PDF = {
    "Times": ("times.ttf", "times.ttf"),
    "TimesBold": ("timesbd.ttf", "timesbd.ttf"),
    "TimesBoldIt": ("timesbi.ttf", "timesbi.ttf"),
    "TimesIt": ("timesi.ttf", "timesi.ttf"),
    "Helvetica": ("Helvetica.ttf", "arial.ttf"),
    "HelveticaBold": ("Helvetica-Bold.ttf", "arialbd.ttf"),
    "HelveticaBoldIt": ("Helvetica-BoldOblique.ttf", "arialbi.ttf"),
    "HelveticaIt": ("Helvetica-Oblique.ttf", "ariali.ttf"),
    "Courier": ("cour.ttf", "cour.ttf"),
    "CourierBold": ("courbd.ttf", "courbd.ttf"),
    "CourierBoldIt": ("courbi.ttf", "courbi.ttf"),
    "CourierIt": ("couri.ttf", "couri.ttf"),
    "Calibri": ("calibri.ttf", "calibri.ttf"),
    "CalibriBold": ("calibrib.ttf", "calibrib.ttf"),
    "CalibriBoldIt": ("calibriz.ttf", "calibriz.ttf"),
    "CalibriIt": ("calibrii.ttf", "calibrii.ttf"),
    "Verdana": ("verdana.ttf", "verdana.ttf"),
    "VerdanaBold": ("verdanab.ttf", "verdanab.ttf"),
    "VerdanaBoldIt": ("verdanaz.ttf", "verdanaz.ttf"),
    "VerdanaIt": ("verdanai.ttf", "verdanai.ttf"),
    "Georgia": ("georgia.ttf", "georgia.ttf"),
    "GeorgiaBold": ("georgiab.ttf", "georgiab.ttf"),
    "GeorgiaBoldIt": ("georgiaz.ttf", "georgiaz.ttf"),
    "GeorgiaIt": ("georgiai.ttf", "georgiai.ttf"),
    "SegoeUI": ("segoeui.ttf", "segoeui.ttf"),
    "SegoeUIBold": ("segoeuib.ttf", "segoeuib.ttf"),
    "SegoeUIBoldIt": ("segoeuiz.ttf", "segoeuiz.ttf"),
    "SegoeUIIt": ("segoeuii.ttf", "segoeuii.ttf"),
}
# Lettera del colore nei codici del .gui -> RGB (g = verde, come v, per i .gui vecchi)
COLORI_TESTO = {
    "n": (0, 0, 0),
    "r": (255, 0, 0),
    "v": (0, 128, 0),
    "g": (0, 128, 0),
    "b": (0, 0, 255),
    "w": (255, 255, 255),
}
# fpdf salva i dati dei font in file .pkl; li mette nella cartella temporanea con un nome
# legato al percorso completo del font, così non si usano .pkl vecchi con percorsi sbagliati
fpdf.fpdf.FPDF_CACHE_MODE = 2
fpdf.fpdf.FPDF_CACHE_DIR = tempfile.gettempdir()
# Solo quando è avviato come programma (non quando viene importato)
if __name__ == "__main__":
    #Crea un array con i file precedenti nella cartella
    filesprec = []

    process(filesprec)

    #loop che avviene ogni secondo
    while True:
        time.sleep(1)
        filesprec = process(filesprec)
