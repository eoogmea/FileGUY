"""Finestra per ritoccare la formattazione di un documento Word prima di salvarlo.

Uso da SingleGUY:
    editor_word.modifica_documento(document, "titolo")
Il documento (python-docx) viene mostrato con la formattazione già applicata dal .gui;
si seleziona il testo e si applicano grassetto, corsivo, tipo e dimensione del carattere.
Con "Salva" le modifiche vengono riscritte nel documento, che poi SingleGUY salva come prima.
"""
import tkinter as tk
from tkinter import messagebox, ttk

from docx.shared import Pt, RGBColor

FONT = ["Times New Roman", "Arial", "Courier New", "Calibri", "Verdana", "Georgia"]
DIMENSIONI = ["8", "9", "10", "11", "12", "14", "16", "18", "20", "24", "28", "36"]
FONT_PREDEFINITO = "Times New Roman"
DIMENSIONE_PREDEFINITA = 11

# Ogni carattere del testo ha dei tag "logici" che descrivono la formattazione:
#   "b" grassetto, "i" corsivo, "fam:<font>", "sz:<dimensione>", "col:<RRGGBB>"
# e un tag "di visualizzazione" (prefisso "v|") che li combina in un font di Tk.


def _nome_font(nome):
    # SingleGUY usa nomi come TimesBold o HelveticaIt, che Word non conosce
    if not nome:
        return FONT_PREDEFINITO
    for prefisso, vero in (("Times", "Times New Roman"), ("Helvetica", "Arial"), ("Courier", "Courier New")):
        if nome.startswith(prefisso):
            return vero
    return nome


def leggi_segmenti(document):
    """Trasforma il documento in una lista di (testo, attributi)."""
    segmenti = []
    for n, par in enumerate(document.paragraphs):
        if n > 0:
            segmenti.append(("\n", {}))
        for run in par.runs:
            f = run.font
            colore = f.color.rgb if f.color.type is not None else None
            segmenti.append((run.text, {
                "fam": _nome_font(f.name),
                "sz": int(round(f.size.pt)) if f.size is not None else DIMENSIONE_PREDEFINITA,
                "b": bool(f.bold),
                "i": bool(f.italic),
                "col": str(colore) if colore is not None and str(colore) != "000000" else None,
            }))
    return segmenti


def scrivi_segmenti(document, segmenti):
    """Sostituisce il contenuto del documento con i segmenti (margini e stili restano)."""
    body = document.element.body
    for par in list(document.paragraphs):
        body.remove(par._element)
    par = document.add_paragraph()
    for testo, attributi in segmenti:
        run = par.add_run(testo)
        font = run.font
        font.name = attributi["fam"]
        font.size = Pt(attributi["sz"])
        font.bold = attributi["b"]
        font.italic = attributi["i"]
        if attributi["col"]:
            font.color.rgb = RGBColor.from_string(attributi["col"])


class EditorWord(tk.Tk):
    def __init__(self, segmenti, titolo=""):
        super().__init__()
        self.title("Editor Word" + (" - " + titolo if titolo else ""))
        self.geometry("900x650")
        self.salvato = False
        self.modificato = False
        self.risultato = []

        self._crea_barra()
        self._crea_testo()
        self.stato = ttk.Label(self, text="Seleziona il testo e scegli la formattazione", foreground="gray",
                               padding=(8, 2))
        self.stato.pack(side="bottom", fill="x")

        for testo, attributi in segmenti:
            self.testo.insert("end", testo, self._tag_da_attributi(attributi))
        self._aggiorna_vista()
        self.testo.edit_reset()
        self.testo.edit_modified(False)
        self.testo.mark_set("insert", "1.0")
        self.testo.focus_set()
        self._aggiorna_barra()
        self.protocol("WM_DELETE_WINDOW", self._chiudi)

    # --- Interfaccia -------------------------------------------------------

    def _crea_barra(self):
        barra = ttk.Frame(self, padding=(6, 6))
        barra.pack(side="top", fill="x")

        self.pulsante_b = tk.Button(barra, text="G", width=3, font=("Arial", 10, "bold"),
                                    command=lambda: self._alterna("b"))
        self.pulsante_b.pack(side="left")
        self.pulsante_i = tk.Button(barra, text="C", width=3, font=("Arial", 10, "italic"),
                                    command=lambda: self._alterna("i"))
        self.pulsante_i.pack(side="left", padx=(2, 10))

        ttk.Label(barra, text="Carattere").pack(side="left")
        self.scelta_font = ttk.Combobox(barra, values=FONT, width=20)
        self.scelta_font.pack(side="left", padx=(4, 10))
        self.scelta_font.bind("<<ComboboxSelected>>", lambda e: self._imposta("fam", self.scelta_font.get()))
        self.scelta_font.bind("<Return>", lambda e: self._imposta("fam", self.scelta_font.get()))

        ttk.Label(barra, text="Dimensione").pack(side="left")
        self.scelta_dimensione = ttk.Combobox(barra, values=DIMENSIONI, width=5)
        self.scelta_dimensione.pack(side="left", padx=(4, 10))
        self.scelta_dimensione.bind("<<ComboboxSelected>>", lambda e: self._imposta_dimensione())
        self.scelta_dimensione.bind("<Return>", lambda e: self._imposta_dimensione())

        ttk.Button(barra, text="Annulla modifiche", command=self._annulla).pack(side="right")
        ttk.Button(barra, text="Salva", command=self._salva).pack(side="right", padx=(0, 6))

    def _crea_testo(self):
        cornice = ttk.Frame(self)
        cornice.pack(side="top", fill="both", expand=True, padx=6)
        scorrimento = ttk.Scrollbar(cornice, orient="vertical")
        scorrimento.pack(side="right", fill="y")
        # Il testo senza tag usa il carattere predefinito, lo stesso con cui verrà salvato
        self.testo = tk.Text(cornice, wrap="word", undo=True, padx=12, pady=12,
                             font=(FONT_PREDEFINITO, DIMENSIONE_PREDEFINITA), yscrollcommand=scorrimento.set)
        self.testo.pack(side="left", fill="both", expand=True)
        scorrimento.configure(command=self.testo.yview)

        self.testo.bind("<Key>", self._digita)
        # Scorciatoie come in Word italiano (Ctrl+G grassetto, Ctrl+I corsivo), più Ctrl+B
        for tasto in ("<Control-g>", "<Control-G>", "<Control-b>", "<Control-B>"):
            self.testo.bind(tasto, lambda e: self._alterna("b") or "break")
        for tasto in ("<Control-i>", "<Control-I>"):
            self.testo.bind(tasto, lambda e: self._alterna("i") or "break")
        for evento in ("<KeyRelease>", "<ButtonRelease-1>", "<<Selection>>"):
            self.testo.bind(evento, lambda e: self._aggiorna_barra(), add="+")
        self.testo.bind("<<Paste>>", lambda e: self.after_idle(self._aggiorna_vista), add="+")
        self.testo.bind("<<Modified>>", self._segna_modificato)

    # --- Tag e formattazione -----------------------------------------------

    def _tag_da_attributi(self, attributi):
        if not attributi:
            return ()
        tag = ["fam:" + attributi["fam"], "sz:" + str(attributi["sz"])]
        if attributi["b"]:
            tag.append("b")
        if attributi["i"]:
            tag.append("i")
        if attributi["col"]:
            tag.append("col:" + attributi["col"])
        return tuple(tag)

    def _attributi_da_tag(self, tags):
        attributi = {"fam": FONT_PREDEFINITO, "sz": DIMENSIONE_PREDEFINITA, "b": False, "i": False, "col": None}
        for tag in tags:
            if tag == "b" or tag == "i":
                attributi[tag] = True
            elif tag.startswith("fam:"):
                attributi["fam"] = tag[4:]
            elif tag.startswith("sz:"):
                attributi["sz"] = int(tag[3:])
            elif tag.startswith("col:"):
                attributi["col"] = tag[4:]
        return attributi

    def _tag_logici(self, indice):
        return [t for t in self.testo.tag_names(indice) if t != "sel"]

    def _scorri_testo(self):
        """Restituisce i pezzi di testo (inizio, testo, tag logici attivi) dall'inizio alla fine."""
        attivi = set()
        pezzi = []
        for chiave, valore, indice in self.testo.dump("1.0", "end-1c", tag=True, text=True):
            if chiave == "tagon":
                attivi.add(valore)
            elif chiave == "tagoff":
                attivi.discard(valore)
            elif chiave == "text":
                pezzi.append((indice, valore, {t for t in attivi if t != "sel" and not t.startswith("v|")}))
        return pezzi

    def _aggiorna_vista(self):
        # Ricalcola i tag di visualizzazione combinando i tag logici di ogni pezzo di testo
        for tag in self.testo.tag_names():
            if tag.startswith("v|"):
                self.testo.tag_remove(tag, "1.0", "end")
        for indice, valore, tags in self._scorri_testo():
            if not tags:
                continue
            a = self._attributi_da_tag(tags)
            stile = " ".join(s for s, attivo in (("bold", a["b"]), ("italic", a["i"])) if attivo) or "normal"
            nome = "v|" + a["fam"] + "|" + str(a["sz"]) + "|" + stile + "|" + (a["col"] or "")
            self.testo.tag_configure(nome, font=(a["fam"], a["sz"], stile), foreground="#" + (a["col"] or "000000"))
            self.testo.tag_add(nome, indice, indice + "+" + str(len(valore)) + "c")
        self.testo.tag_raise("sel")

    def _selezione(self):
        try:
            return self.testo.index("sel.first"), self.testo.index("sel.last")
        except tk.TclError:
            self._messaggio("Seleziona prima il testo da formattare")
            return None

    def _tutto_con(self, tag, inizio, fine):
        indice = inizio
        while self.testo.compare(indice, "<", fine):
            if tag not in self.testo.tag_names(indice):
                return False
            indice = self.testo.index(indice + "+1c")
        return True

    def _alterna(self, tag):
        selezione = self._selezione()
        if selezione is None:
            return
        inizio, fine = selezione
        if self._tutto_con(tag, inizio, fine):
            self.testo.tag_remove(tag, inizio, fine)
        else:
            self.testo.tag_add(tag, inizio, fine)
        self._dopo_formattazione()

    def _imposta(self, prefisso, valore):
        valore = valore.strip()
        selezione = self._selezione()
        if selezione is None or valore == "":
            return
        inizio, fine = selezione
        for tag in self.testo.tag_names():
            if tag.startswith(prefisso + ":"):
                self.testo.tag_remove(tag, inizio, fine)
        self.testo.tag_add(prefisso + ":" + valore, inizio, fine)
        self._dopo_formattazione()

    def _imposta_dimensione(self):
        valore = self.scelta_dimensione.get().strip()
        if not valore.isdigit() or not 1 <= int(valore) <= 400:
            self._messaggio("Dimensione non valida: " + valore)
            return
        self._imposta("sz", str(int(valore)))

    def _dopo_formattazione(self):
        self.modificato = True
        self._aggiorna_vista()
        self._aggiorna_barra()
        self.testo.focus_set()

    def _digita(self, evento):
        # Il testo scritto prende la formattazione del carattere che lo precede, come in Word
        ctrl = evento.state & 0x4 and not evento.state & 0x20000  # AltGr = Ctrl+Alt, va lasciato passare
        if ctrl or evento.char == "":
            return None
        if evento.char == "\r":
            carattere = "\n"
        elif evento.char == "\t" or evento.char >= " " and evento.char != "\x7f":
            carattere = evento.char
        else:
            return None
        selezione = self.testo.tag_ranges("sel")
        if selezione:
            tags = self._tag_logici(selezione[0])
            self.testo.delete(selezione[0], selezione[1])
        elif self.testo.compare("insert", ">", "1.0"):
            tags = self._tag_logici("insert-1c")
        else:
            tags = self._tag_logici("insert")
        self.testo.insert("insert", carattere, tuple(tags))
        self.testo.see("insert")
        return "break"

    def _aggiorna_barra(self):
        selezione = self.testo.tag_ranges("sel")
        if selezione:
            indice = selezione[0]
        elif self.testo.compare("insert", ">", "1.0"):
            indice = "insert-1c"
        else:
            indice = "insert"
        a = self._attributi_da_tag(self._tag_logici(indice))
        self.scelta_font.set(a["fam"])
        self.scelta_dimensione.set(str(a["sz"]))
        self.pulsante_b.configure(relief="sunken" if a["b"] else "raised")
        self.pulsante_i.configure(relief="sunken" if a["i"] else "raised")

    # --- Risultato ---------------------------------------------------------

    def segmenti(self):
        """Il testo con la formattazione, unendo i pezzi consecutivi uguali."""
        risultato = []
        for _, valore, tags in self._scorri_testo():
            attributi = self._attributi_da_tag(tags)
            if risultato and risultato[-1][1] == attributi:
                risultato[-1] = (risultato[-1][0] + valore, attributi)
            else:
                risultato.append((valore, attributi))
        return risultato

    def _segna_modificato(self, evento):
        if self.testo.edit_modified():
            self.modificato = True

    def _messaggio(self, testo):
        self.stato.configure(text=testo)

    def _termina(self, salva):
        # Il testo va letto prima di distruggere la finestra
        self.salvato = salva
        if salva:
            self.risultato = self.segmenti()
        self.destroy()

    def _salva(self):
        self._termina(True)

    def _annulla(self):
        if self.modificato and not messagebox.askyesno(
                "Annulla", "Le modifiche andranno perse e verrà salvato il documento originale. Continuare?",
                parent=self):
            return
        self._termina(False)

    def _chiudi(self):
        if not self.modificato:
            self._termina(False)
            return
        risposta = messagebox.askyesnocancel("Chiudi", "Salvare le modifiche al documento?", parent=self)
        if risposta is None:
            return
        self._termina(risposta)


def modifica_documento(document, titolo=""):
    """Apre l'editor sul documento; se l'utente salva, riscrive il documento. Restituisce True se modificato."""
    editor = EditorWord(leggi_segmenti(document), titolo)
    editor.mainloop()
    if editor.salvato:
        scrivi_segmenti(document, editor.risultato)
        return True
    return False


if __name__ == "__main__":
    # Prova: python editor_word.py documento.docx  ->  salva documento_modificato.docx
    import os
    import sys

    from docx import Document

    if len(sys.argv) < 2:
        print("Uso: python editor_word.py documento.docx")
        sys.exit(1)
    documento = Document(sys.argv[1])
    if modifica_documento(documento, os.path.basename(sys.argv[1])):
        uscita = os.path.splitext(sys.argv[1])[0] + "_modificato.docx"
        documento.save(uscita)
        print("Salvato " + uscita)
