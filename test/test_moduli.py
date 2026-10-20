"""Prova dei moduli di supporto: stampa_pdf, editor_word, config_sicuro, whatsapp_meta.

Uso:  python test\\test_moduli.py
Non stampa su carta: usa la stampante virtuale "Microsoft Print to PDF" con un file di uscita.
Non invia messaggi: le richieste a WhatsApp sono intercettate e solo controllate.
"""
import json
import os
import shutil
import sys
import tempfile
import traceback

import fitz

CARTELLA_PROGRAMMA = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, CARTELLA_PROGRAMMA)

risultati = []


def prova(funzione):
    try:
        funzione()
        risultati.append((funzione.__name__, None))
    except Exception:
        risultati.append((funzione.__name__, traceback.format_exc(limit=-3).strip()))
    return funzione


cartella = tempfile.mkdtemp(prefix="infas_test_")


@prova
def stampa_su_microsoft_print_to_pdf():
    from stampa_pdf import stampa_pdf
    sorgente = os.path.join(cartella, "da_stampare.pdf")
    doc = fitz.open()
    for n in range(2):
        pagina = doc.new_page(width=595, height=842)
        pagina.insert_text((72, 100), "Pagina di prova %d" % (n + 1), fontsize=20)
    doc.save(sorgente)
    doc.close()
    uscita = os.path.join(cartella, "stampato.pdf")
    stampa_pdf(sorgente, "Microsoft Print to PDF", uscita)
    assert os.path.isfile(uscita), "la stampante virtuale non ha creato il file"
    stampato = fitz.open(uscita)
    assert stampato.page_count == 2, "pagine stampate: %d" % stampato.page_count
    stampato.close()


@prova
def editor_word_legge_e_riscrive():
    from docx import Document
    from docx.shared import Pt, RGBColor
    import editor_word
    doc = Document()
    par = doc.add_paragraph()
    run = par.add_run("Titolo\n")
    run.font.name, run.font.size, run.font.bold = "TimesBold", Pt(14), True
    run = par.add_run("testo rosso")
    run.font.name, run.font.size, run.font.italic = "HelveticaIt", Pt(10), True
    run.font.color.rgb = RGBColor(255, 0, 0)
    segmenti = editor_word.leggi_segmenti(doc)
    assert segmenti[0][1]["fam"] == "Times New Roman" and segmenti[0][1]["b"], segmenti[0]
    assert segmenti[1][1]["fam"] == "Arial" and segmenti[1][1]["col"] == "FF0000", segmenti[1]
    editor_word.scrivi_segmenti(doc, segmenti)
    percorso = os.path.join(cartella, "editor.docx")
    doc.save(percorso)
    riletto = Document(percorso)
    assert "".join(p.text for p in riletto.paragraphs) == "Titolo\ntesto rosso"


@prova
def editor_word_finestra():
    # Apre la finestra e la chiude subito con "Salva" senza modifiche
    from docx import Document
    import editor_word
    doc = Document()
    doc.add_paragraph().add_run("prova finestra")
    originale = editor_word.EditorWord.mainloop

    def mainloop(self, n=0):
        self.after(300, lambda: (setattr(self, "salvato", True),
                                 setattr(self, "risultato", editor_word.leggi_segmenti(doc)), self.destroy()))
        originale(self, n)
    editor_word.EditorWord.mainloop = mainloop
    try:
        assert editor_word.modifica_documento(doc, "prova") is True
    finally:
        editor_word.EditorWord.mainloop = originale
    assert "".join(p.text for p in doc.paragraphs) == "prova finestra"


@prova
def config_cifrata_salva_e_rilegge():
    import config_sicuro
    percorso = os.path.join(cartella, "config.dat")
    dati = dict(config_sicuro.CAMPI, cartella_lavoro="C:\\lavoro", email_password="segreta è")
    config_sicuro.salva_config(dati, percorso)
    with open(percorso, "rb") as f:
        assert b"segreta" not in f.read(), "la password è in chiaro nel file"
    assert config_sicuro.carica_config(percorso) == dati
    config_sicuro.salva_config(dati, percorso, tutto_il_pc=True)
    assert config_sicuro.carica_config(percorso) == dati


@prova
def config_vera_leggibile():
    import config_sicuro
    if config_sicuro.config_esiste():
        config = config_sicuro.carica_config()
        for campo in ("cartella_lavoro", "cartella_template", "cartella_output"):
            if config[campo] and not os.path.isdir(config[campo]):
                raise AssertionError("config.dat: la cartella %s non esiste: %s" % (campo, config[campo]))


@prova
def whatsapp_numeri_e_richieste():
    import whatsapp_meta
    assert whatsapp_meta.normalizza_numero("+39 333 123.45-67") == "393331234567"
    assert whatsapp_meta.normalizza_numero("0039 3331234567") == "393331234567"
    try:
        whatsapp_meta.normalizza_numero("333abc")
        raise AssertionError("numero non valido accettato")
    except whatsapp_meta.WhatsAppError:
        pass

    inviati = []

    class Risposta:
        ok, status_code, text = True, 200, ""

        def __init__(self, dati):
            self.dati = dati

        def json(self):
            return self.dati

    wa = whatsapp_meta.WhatsApp("token", "12345")

    def post(url, json=None, data=None, files=None, timeout=None):
        inviati.append((url, json, data))
        if url.endswith("/media"):
            return Risposta({"id": "media-1"})
        return Risposta({"messages": [{"id": "wamid.1"}]})
    wa.session.post = post

    allegato = os.path.join(cartella, "fattura.pdf")
    with open(allegato, "wb") as f:
        f.write(b"%PDF-1.4\n")
    assert wa.invia_testo("+39 333 1234567", "ciao") == "wamid.1"
    wa.invia_documento("393331234567", allegato)
    wa.invia_template("393331234567", "fattura", "it", ["Mario", "123"], allegato)
    testo, media, documento, media2, template = inviati
    assert testo[0].endswith("/v23.0/12345/messages") and testo[1]["to"] == "393331234567"
    assert documento[1]["document"] == {"id": "media-1", "filename": "fattura.pdf"}
    componenti = template[1]["template"]["components"]
    assert componenti[0]["type"] == "header" and componenti[1]["parameters"][1]["text"] == "123"
    assert wa.session.headers["Authorization"] == "Bearer token"

    class Errore(Risposta):
        ok, status_code = False, 401
    wa.session.post = lambda *a, **k: Errore({"error": {"message": "Invalid OAuth", "code": 190}})
    try:
        wa.invia_testo("393331234567", "x")
        raise AssertionError("errore API non segnalato")
    except whatsapp_meta.WhatsAppError as e:
        assert "190" in str(e)


@prova
def config_manager_si_avvia():
    import ConfigManager
    finestra = ConfigManager.ConfigManager()
    finestra.update()
    finestra.destroy()


shutil.rmtree(cartella, ignore_errors=True)
for nome, errore in risultati:
    print("[%s] %s" % ("OK " if errore is None else "ERR", nome))
    if errore:
        print("       " + errore.replace("\n", "\n       "))
print("\n%d/%d prove OK" % (sum(e is None for _, e in risultati), len(risultati)))
sys.exit(0 if all(e is None for _, e in risultati) else 1)
