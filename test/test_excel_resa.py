"""Prova completa dei formati Excel, pensata per Linux (funziona anche su Windows).

Uso:  python test/test_excel_resa.py                 prova e confronto con il riferimento
      python test/test_excel_resa.py --riferimento   (su Windows) riscrive il riferimento

1. Il servizio vero (SingleGUY.main, niente simulazioni) crea un .xlsx con tutti i 25 formati
   numerici, tutti i font, grassetto/corsivo, i colori, gli sfondi, gli allineamenti, le larghezze
   e le righe speciali (intestazione testo=si, ultima riga u1).
2. Il file si rilegge con openpyxl: ogni cella deve avere valore, formato, carattere, colore,
   sfondo e allineamento attesi.
3. Tutto il file (cella per cella) si confronta con test/casi/excel_resa_riferimento.json, creato
   su Windows: una differenza vuol dire che su questo sistema il servizio scrive un file diverso.
4. Se c'è LibreOffice (soffice) il file si apre davvero: il testo delle celle come lo mostra Calc
   (es. "1.234,56 €", "20-ott-2026") si confronta con quello atteso in italiano, e il foglio diventa
   un PDF e un'immagine da guardare (test/_run/excel_resa/excel_resa.png).
5. Quali font usa il sistema al posto di Calibri, Arial, Times New Roman... (fc-match): con un
   sostituto di misure diverse le colonne e i testi non vengono come su Windows.
"""
import datetime
import json
import os
import shutil
import subprocess
import sys
import tempfile
import traceback

CARTELLA_PROGRAMMA = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CARTELLA_RUN = os.path.join(CARTELLA_PROGRAMMA, "test", "_run", "excel_resa")
RIFERIMENTO = os.path.join(CARTELLA_PROGRAMMA, "test", "casi", "excel_resa_riferimento.json")
sys.path.insert(0, CARTELLA_PROGRAMMA)
os.chdir(CARTELLA_PROGRAMMA)

from openpyxl import load_workbook

import modello_gui

NOME = "excel_resa"

# Font del .gui -> nome che deve finire nel file Excel
FONT_EXCEL = {"t": "Times New Roman", "h": "Arial", "c": "Courier New", "k": "Calibri", "v": "Verdana",
              "g": "Georgia", "u": "Segoe UI", "o": "Consolas", "d": "DejaVu Sans Mono", "p": "IBM Plex Mono",
              "l": "Liberation Mono"}
COLORI_ARGB = {"n": "FF000000", "r": "FFFF0000", "v": "FF008000", "b": "FF0000FF", "w": "FFFFFFFF"}

# Per ogni formato: (testo nella riga 2, testo nella riga 3, come li mostra Calc in italiano).
# None = il testo resta testo (non è un numero per quel formato).
CAMPIONI = {
    0: ("testo libero", "00123", ("testo libero", "00123")),
    1: ("1234", "-7", ("1234", "-7")),
    2: ("1234,56", "0,5", ("1234,56", "0,50")),
    3: ("1234", "1.000.000", ("1.234", "1.000.000")),
    4: ("1.234,56", "-0,5", ("1.234,56", "-0,50")),
    5: ("1234", "-1234", ("1.234 €", "-1.234 €")),
    6: ("1234", "-1234", ("1.234 €", "-1.234 €")),
    7: ("1234", "-1234", ("1.234 €", "-1.234 €")),
    8: ("1.234,56", "-1.234,56", ("1.234,56 €", "-1.234,56 €")),
    9: ("12%", "0,25", ("12%", "25%")),
    10: ("12,5%", "0,125", ("12,50%", "12,50%")),
    11: ("1234", "0,000123", ("1,23E+03", "1,23E-04")),
    12: ("20/10/2026", "01/02/2027", ("20/10/2026", "01/02/2027")),
    13: ("20/10/2026", "01/02/2027", ("20-ott-2026", "01-feb-2027")),
    14: ("20/10/2026", "01/02/2027", ("20-ott", "01-feb")),
    15: ("20/10/2026", "01/02/2027", ("ott 2026", "feb 2027")),
    16: ("13:30", "09:05", ("1:30 PM", "9:05 AM")),
    17: ("13:30:55", "09:05:07", ("1:30:55 PM", "9:05:07 AM")),
    18: ("13:30", "09:05", ("13:30", "09:05")),
    19: ("13:30:55", "09:05:07", ("13:30:55", "09:05:07")),
    20: ("20/10/2026 13:30", "01/02/2027 09:05", ("20/10/2026 13:30", "01/02/2027 09:05")),
    21: ("1234", "0,5", ("1,2E+3", "500,0E-3")),
    22: ("13:30:55", "09:05:07", ("30:55", "05:07")),
    23: ("00123", "1234", ("00123", "1234")),
    24: ("1.234,56", "-1.234,56", ("1.234,56 €", "-1.234,56 €")),
}
CODICI = sorted(CAMPIONI)
LETTERE_FONT = list(FONT_EXCEL)
COLORI = ["n", "r", "v", "b", "n"]
ALLINEAMENTI = ["", "sinistra", "centro", "destra"]


def colonne_del_modello():
    """Una colonna per ogni formato, con font, stile, colore e allineamento che cambiano."""
    colonne = []
    for n, codice in enumerate(CODICI):
        lettera = LETTERE_FONT[n % len(LETTERE_FONT)]
        formato = (lettera, n % 2 == 1, n % 3 == 2, 9 + n % 6, COLORI[n % len(COLORI)])
        colonne.append({"formato": codice, "carattere": formato,
                        "larghezza": 14 + n % 5 if n % 4 == 0 else None,
                        "allineamento": ALLINEAMENTI[n % len(ALLINEAMENTI)] or None})
    return colonne


STILI = [
    {"righe": "1", "nome": "Intestazione", "carattere": ("h", True, False, 11, "w"), "sfondo": None,
     "allineamento": "centro", "testo": True},
    {"righe": "4", "nome": "Evidenziata", "carattere": None, "sfondo": "azzurro", "allineamento": None,
     "testo": False},
    {"righe": "u1", "nome": "Totale", "carattere": ("t", True, True, 12, "r"), "sfondo": "giallo",
     "allineamento": "destra", "testo": False},
]


def righe_del_testo():
    intestazione = ["F%d" % c for c in CODICI]
    riga2 = [CAMPIONI[c][0] for c in CODICI]
    riga3 = [CAMPIONI[c][1] for c in CODICI]
    # Riga 4 (sfondo azzurro) e riga 5 (totale) con gli stessi valori della riga 2
    return [intestazione, riga2, riga3, riga2, riga2]


def crea_excel():
    """Crea il .xlsx con il servizio vero, come farebbe con i file della cartella di lavoro."""
    shutil.rmtree(CARTELLA_RUN, ignore_errors=True)
    lavoro, uscita = (os.path.join(CARTELLA_RUN, c) for c in ("lavoro", "out"))
    os.makedirs(lavoro)
    os.makedirs(uscita)
    with open(os.path.join(lavoro, NOME + ".txt"), "w", encoding="utf-8") as f:
        f.write("".join(";".join(r) + ";\n" for r in righe_del_testo()))
    intestazione = modello_gui.nuovo_modello(False)
    intestazione.update(testo=NOME + ".txt", azione="excel", uscita=NOME + ".xlsx")
    modello_gui.scrivi_gui(os.path.join(lavoro, NOME + ".gui"), intestazione,
                           modello_gui.linee_excel(STILI, colonne_del_modello()))
    import registro
    import SingleGUY
    registro.configura(os.path.join(CARTELLA_RUN, "FileGUY.log"))
    SingleGUY.main(NOME + ".txt", {"cartella_lavoro": lavoro, "cartella_output": uscita, "cartella_template": ""})
    percorso = os.path.join(uscita, NOME + ".xlsx")
    if not os.path.isfile(percorso):
        with open(os.path.join(CARTELLA_RUN, "FileGUY.log"), encoding="utf-8") as f:
            raise AssertionError("il servizio non ha creato l'Excel:\n" + f.read())
    return percorso


def descrivi_cella(cella):
    colore = getattr(cella.font.color, "rgb", None)
    valore = cella.value
    if isinstance(valore, (datetime.datetime, datetime.date, datetime.time)):
        valore = type(valore).__name__ + " " + valore.isoformat()
    return {
        "valore": valore,
        "tipo": type(cella.value).__name__,
        "formato": cella.number_format,
        "font": cella.font.name,
        "dimensione": cella.font.size,
        "grassetto": bool(cella.font.bold),
        "corsivo": bool(cella.font.italic),
        "colore": colore if isinstance(colore, str) else None,
        "sfondo": cella.fill.start_color.rgb if cella.fill.fill_type == "solid" else None,
        "allineamento": cella.alignment.horizontal,
    }


def descrivi_file(percorso):
    ws = load_workbook(percorso).worksheets[0]
    celle = {"%s%d" % (cella.column_letter, cella.row): descrivi_cella(cella) for riga in ws.iter_rows()
             for cella in riga}
    larghezze = {lettera: d.width for lettera, d in sorted(ws.column_dimensions.items())}
    return {"foglio": ws.title, "righe": ws.max_row, "colonne": ws.max_column, "celle": celle,
            "larghezze": larghezze}


def controlla_celle(descrizione):
    """Ogni cella come la vuole il .gui (calcolato qui, non con il codice del servizio)."""
    errori = []
    colonne = colonne_del_modello()
    stili = {1: STILI[0], 4: STILI[1], 5: STILI[2]}
    righe = righe_del_testo()
    if (descrizione["righe"], descrizione["colonne"]) != (len(righe), len(CODICI)):
        errori.append("foglio di %d righe e %d colonne invece di %d e %d" % (
            descrizione["righe"], descrizione["colonne"], len(righe), len(CODICI)))
    for n, colonna in enumerate(colonne):
        lettera = descrizione_lettera(n)
        for r in range(1, len(righe) + 1):
            cella = descrizione["celle"].get("%s%d" % (lettera, r))
            if cella is None:
                errori.append("manca la cella %s%d" % (lettera, r))
                continue
            stile = stili.get(r)
            carattere = stile["carattere"] if stile and stile["carattere"] else colonna["carattere"]
            f, b, i, s, c = carattere
            atteso = {"font": FONT_EXCEL[f], "dimensione": s, "grassetto": b, "corsivo": i,
                      "colore": COLORI_ARGB[c],
                      "sfondo": "FF" + modello_gui.SFONDI[stile["sfondo"]][1] if stile and stile["sfondo"] else None,
                      "allineamento": {"sinistra": "left", "centro": "center", "destra": "right"}.get(
                          (stile["allineamento"] if stile and stile["allineamento"] else None)
                          or colonna["allineamento"])}
            testo = stile is not None and stile["testo"]
            formato = modello_gui.FORMATI_EXCEL[colonna["formato"]][1]
            atteso["formato"] = "General" if testo or not formato else formato
            for chiave, valore in atteso.items():
                if cella[chiave] != valore:
                    errori.append("%s%d (formato %d) %s: %r invece di %r" % (lettera, r, colonna["formato"], chiave,
                                                                            cella[chiave], valore))
            # Nelle righe di dati i valori dei formati numerici, data e ora non devono restare testo
            if not testo and colonna["formato"] not in (0, 23) and cella["tipo"] == "str":
                errori.append("%s%d (formato %d): %r è rimasto testo" % (lettera, r, colonna["formato"],
                                                                       cella["valore"]))
        if colonna["larghezza"] and descrizione["larghezze"].get(lettera) != colonna["larghezza"]:
            errori.append("colonna %s larga %r invece di %r" % (lettera, descrizione["larghezze"].get(lettera),
                                                               colonna["larghezza"]))
    return errori


def descrizione_lettera(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n + 1)


def confronta_con_riferimento(descrizione):
    if not os.path.isfile(RIFERIMENTO):
        return ["manca il riferimento %s: crealo su Windows con --riferimento" % RIFERIMENTO]
    with open(RIFERIMENTO, encoding="utf-8") as f:
        riferimento = json.load(f)
    errori = []
    for chiave in ("foglio", "righe", "colonne"):
        if descrizione[chiave] != riferimento[chiave]:
            errori.append("%s: %r invece di %r (Windows)" % (chiave, descrizione[chiave], riferimento[chiave]))
    for nome in sorted(set(descrizione["celle"]) | set(riferimento["celle"])):
        qui, la = descrizione["celle"].get(nome), riferimento["celle"].get(nome)
        if qui != la:
            diverse = sorted(k for k in set(qui or {}) | set(la or {}) if (qui or {}).get(k) != (la or {}).get(k))
            errori.append("cella %s: %s" % (nome, ", ".join("%s %r invece di %r" % (k, (qui or {}).get(k),
                                                                                      (la or {}).get(k))
                                                             for k in diverse)))
    for lettera in sorted(set(descrizione["larghezze"]) | set(riferimento["larghezze"])):
        if descrizione["larghezze"].get(lettera) != riferimento["larghezze"].get(lettera):
            errori.append("colonna %s larga %r invece di %r (Windows)" % (
                lettera, descrizione["larghezze"].get(lettera), riferimento["larghezze"].get(lettera)))
    return errori


def trova_soffice():
    for nome in ("soffice", "libreoffice"):
        if shutil.which(nome):
            return shutil.which(nome)
    for percorso in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                     r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                     "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if os.path.isfile(percorso):
            return percorso
    return None


def converti(soffice, percorso, filtro, cartella):
    """Converte con LibreOffice senza finestre, con un profilo nuovo e la lingua italiana."""
    profilo = tempfile.mkdtemp(prefix="lo_profilo_")
    ambiente = dict(os.environ, LANG="it_IT.UTF-8", LC_ALL="it_IT.UTF-8", LANGUAGE="it")
    try:
        comando = [soffice, "-env:UserInstallation=" + _url_file(profilo), "--headless", "--norestore",
                   "--convert-to", filtro, "--outdir", cartella, percorso]
        esito = subprocess.run(comando, capture_output=True, text=True, timeout=180, env=ambiente)
        if esito.returncode != 0:
            raise AssertionError("LibreOffice non ha convertito il file (%s): %s %s" % (
                filtro, esito.stdout.strip(), esito.stderr.strip()))
    finally:
        shutil.rmtree(profilo, ignore_errors=True)


def _url_file(percorso):
    percorso = os.path.abspath(percorso).replace("\\", "/")
    return "file://" + ("/" if not percorso.startswith("/") else "") + percorso


def controlla_libreoffice(percorso):
    """Il testo delle celle come lo mostra Calc e un'immagine del foglio."""
    soffice = trova_soffice()
    if soffice is None:
        return None, ["LibreOffice non installato: il controllo di come si vede il file è saltato "
                      "(su Debian/Ubuntu: sudo apt install libreoffice-calc)"]
    note = []
    cartella = os.path.join(CARTELLA_RUN, "libreoffice")
    os.makedirs(cartella, exist_ok=True)
    # 59 = ";", 34 = ", 76 = UTF-8, prima riga 1, ..., true finale = "contenuto delle celle come mostrato"
    converti(soffice, percorso, "csv:Text - txt - csv (StarCalc):59,34,76,1,,1040,false,true,true", cartella)
    import csv
    with open(os.path.join(cartella, NOME + ".csv"), encoding="utf-8") as f:
        mostrate = list(csv.reader(f, delimiter=";"))
    errori = []
    for n, codice in enumerate(CODICI):
        for r, atteso in ((2, CAMPIONI[codice][2][0]), (3, CAMPIONI[codice][2][1])):
            trovato = mostrate[r - 1][n] if r - 1 < len(mostrate) and n < len(mostrate[r - 1]) else None
            if trovato != atteso:
                errori.append("%s%d (formato %d, %s): Calc mostra %r invece di %r" % (
                    descrizione_lettera(n), r, codice, modello_gui.FORMATI_EXCEL[codice][0].split("  ")[0],
                    trovato, atteso))
    if mostrate and mostrate[0][:1] != ["F0"]:
        errori.append("intestazione mostrata come %r invece di testo F0, F1..." % mostrate[0][:3])
    # Il foglio come lo stampa Calc, da guardare (font, colori, sfondi, allineamenti)
    try:
        converti(soffice, percorso, "pdf", cartella)
        import pymupdf
        with pymupdf.open(os.path.join(cartella, NOME + ".pdf")) as doc:
            for p, pagina in enumerate(doc):
                immagine = os.path.join(CARTELLA_RUN, "%s%s.png" % (NOME, "" if p == 0 else "_%d" % (p + 1)))
                pagina.get_pixmap(dpi=110).save(immagine)
                note.append("immagine da guardare: " + immagine)
    except Exception as e:
        note.append("immagine del foglio non creata: %s" % e)
    return errori, note


def font_sostituiti():
    """Per ogni font del .gui quale usa il sistema (solo dove c'è fontconfig)."""
    if not shutil.which("fc-match"):
        return []
    # Sostituti con le stesse misure: il testo occupa lo stesso spazio che su Windows
    uguali = {"Calibri": "Carlito", "Arial": "Liberation Sans", "Times New Roman": "Liberation Serif",
              "Courier New": "Liberation Mono", "Georgia": "Gelasio"}
    note = []
    for nome in sorted(set(FONT_EXCEL.values())):
        esito = subprocess.run(["fc-match", "-f", "%{family}", nome], capture_output=True, text=True)
        trovato = esito.stdout.split(",")[0].strip()
        if trovato.lower() == nome.lower():
            continue
        if uguali.get(nome, "").lower() == trovato.lower():
            note.append("%s -> %s (stesse misure, va bene)" % (nome, trovato))
        else:
            note.append("%s -> %s: misure diverse, colonne e testi possono venire diversi da Windows%s" % (
                nome, trovato, " (installa %s)" % uguali[nome] if nome in uguali else ""))
    return note


def main(argomenti):
    risultati = []
    note = []

    def prova(nome, funzione, *a):
        try:
            errori = funzione(*a)
        except Exception:
            errori = [traceback.format_exc(limit=-3).strip()]
        risultati.append((nome, errori))
        return errori

    percorso = None
    try:
        percorso = crea_excel()
        risultati.append(("servizio_crea_excel", []))
    except Exception:
        risultati.append(("servizio_crea_excel", [traceback.format_exc(limit=-3).strip()]))
    if percorso:
        descrizione = descrivi_file(percorso)
        with open(os.path.join(CARTELLA_RUN, "stili_%s.json" % sys.platform), "w", encoding="utf-8") as f:
            json.dump(descrizione, f, ensure_ascii=False, indent=1, default=str)
        if "--riferimento" in argomenti:
            with open(RIFERIMENTO, "w", encoding="utf-8") as f:
                json.dump(descrizione, f, ensure_ascii=False, indent=1, default=str)
            note.append("riferimento scritto: " + RIFERIMENTO)
        prova("celle_come_nel_gui", controlla_celle, descrizione)
        # json.dump/load: i valori si confrontano come nel file di riferimento
        prova("uguale_al_file_creato_su_windows", confronta_con_riferimento,
              json.loads(json.dumps(descrizione, default=str)))
        try:
            errori, note_lo = controlla_libreoffice(percorso)
        except Exception:
            errori, note_lo = [traceback.format_exc(limit=-3).strip()], []
        if errori is None:
            note += note_lo
        else:
            risultati.append(("come_lo_mostra_libreoffice", errori))
            note += note_lo
    note += font_sostituiti()

    for nome, errori in risultati:
        print("[%s] %s" % ("OK " if not errori else "ERR", nome))
        for e in errori[:60]:
            print("       " + e.replace("\n", "\n       "))
        if len(errori) > 60:
            print("       ... altri %d" % (len(errori) - 60))
    for n in note:
        print("  nota: " + n)
    print("\n%d/%d prove OK (formati Excel, %s)" % (sum(not e for _, e in risultati), len(risultati), sys.platform))
    print("File creato: %s" % percorso)
    return 0 if all(not e for _, e in risultati) else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
