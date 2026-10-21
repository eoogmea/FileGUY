"""Prova dell'editor dei modelli Excel (EditorExcel.py) e delle linee Excel del .gui.

Uso:  python test\\test_editor_excel.py

1. Righe speciali e colonne -> linee -> righe speciali e colonne: deve tornare uguale.
2. Gli intervalli di righe (1-3, 10, u1...) e i .gui di prima, senza la linea "colonna".
3. Le celle mostrate nell'editor come le mostrerà Excel.
4. La finestra: colonne (formato, carattere, larghezza, allineamento), righe speciali (intestazione,
   totale sull'ultima riga, righe scelte con un clic), annulla, copia formato, salvataggio,
   riapertura, anteprima con il codice del servizio e "Applica".
"""
import datetime
import os
import random
import shutil
import sys
import tempfile
import traceback

CARTELLA_PROGRAMMA = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, CARTELLA_PROGRAMMA)
os.chdir(CARTELLA_PROGRAMMA)
import modello_gui

risultati = []
cartella = tempfile.mkdtemp(prefix="infas_excel_")

TESTO = ("Codice;Descrizione;Quantità;Prezzo;Data;Note;\n"
         "1;Viti;100;0,15;20/10/2026;;\n"
         "2;Bulloni;50;1.234,50;21/10/2026;urgente;\n"
         "3;Rondelle più grandi;7;-12;non so;;\n"
         "Totale;;157;1.222,65;;;\n")


def prova(funzione):
    try:
        funzione()
        risultati.append((funzione.__name__, None))
    except Exception:
        risultati.append((funzione.__name__, traceback.format_exc(limit=-3).strip()))
    return funzione


@prova
def linee_andata_e_ritorno():
    casuale = random.Random(7)
    fonts, colori = list(modello_gui.FONT), list(modello_gui.COLORI)

    def carattere():
        if casuale.random() < 0.3:
            return None
        return (casuale.choice(fonts), casuale.random() < 0.5, casuale.random() < 0.5, casuale.randint(6, 36),
                casuale.choice(colori))
    for _ in range(200):
        colonne = [{"formato": casuale.choice(list(modello_gui.FORMATI_EXCEL)), "carattere": carattere(),
                    "larghezza": casuale.choice([None, 1, 12, 255]),
                    "allineamento": casuale.choice([None] + list(modello_gui.ALLINEAMENTI))}
                   for _ in range(casuale.randint(0, 5))]
        # L'ultima colonna non deve essere vuota, altrimenti non si scrive (ed è giusto così)
        colonne.append(dict(modello_gui.colonna_excel_vuota(), formato=4))
        stili = [{"righe": casuale.choice(["1", "1-3, 10", "u1", "2-u2", ""]),
                  "nome": casuale.choice(["", "Totale", "Titolo del listino"]), "carattere": carattere(),
                  "sfondo": casuale.choice([None] + list(modello_gui.SFONDI)),
                  "allineamento": casuale.choice([None] + list(modello_gui.ALLINEAMENTI)),
                  "testo": casuale.random() < 0.5}
                 for _ in range(casuale.randint(0, 3))]
        linee = modello_gui.linee_excel(stili, colonne)
        assert modello_gui.leggi_excel(linee) == (stili, colonne), linee


@prova
def intervalli_di_righe():
    righe = modello_gui.righe_excel
    assert righe("1-3, 10, u1", 20) == {1, 2, 3, 10, 20}
    assert righe("u3-u1", 10) == {8, 9, 10} and righe("u", 10) == {10} and righe("3-1", 10) == {1, 2, 3}
    assert righe("2-u2", 6) == {2, 3, 4, 5} and righe("", 5) == set()
    # Le righe che il foglio non ha si ignorano
    assert righe("4-9, u7", 5) == {4, 5}
    for sbagliato in ("a", "0", "u0", "1-x", "-3"):
        try:
            righe(sbagliato, 10)
            raise AssertionError("accettato: " + sbagliato)
        except ValueError:
            pass
    import EditorExcel
    assert EditorExcel.intervalli({1, 2, 3, 7, 9, 10}) == "1-3, 7, 9-10"


@prova
def gui_di_prima_e_linee_sbagliate():
    vuota = modello_gui.colonna_excel_vuota()
    # Senza la linea "colonna": tutte colonne, e la prima riga resta testo come allora
    stili, colonne = modello_gui.leggi_excel(["8\n", "12\n", "fine\n", "cliente@example.com"])
    assert colonne == [dict(vuota, formato=8), dict(vuota, formato=12)], colonne
    assert len(stili) == 1 and stili[0]["righe"] == "1" and stili[0]["testo"], stili
    leggi = modello_gui.leggi_colonna_excel
    assert leggi("") == vuota and leggi("x") == vuota and leggi("99") == vuota
    assert leggi("4\tlarghezza=0\tallineamento=alto\tsconosciuta=1") == dict(vuota, formato=4)
    stile = modello_gui.leggi_stile_riga("1-2\tnome=Titolo del listino\tsfondo=viola\ttesto=no")
    assert stile["nome"] == "Titolo del listino" and stile["sfondo"] is None and not stile["testo"]
    # Le ultime colonne senza formato non si scrivono
    assert modello_gui.linee_excel([], [dict(vuota, formato=1), vuota]) == ["colonna", "1", "fine"]
    assert modello_gui.campi_excel("a;b;;c\n") == ["a", "b", ""]


@prova
def celle_come_in_excel():
    import EditorExcel
    mostra = EditorExcel.testo_mostrato
    attesi = [("1234,5", 1, "1235"), ("1234,5", 2, "1234,50"), ("1234567", 3, "1.234.567"),
              ("1.234,5", 4, "1.234,50"), ("1234", 5, "1.234 €"), ("-1234,567", 8, "-1.234,57 €"),
              ("12,5%", 10, "12,50%"), ("20/10/2026", 12, "20/10/2026"), ("20/10/2026", 13, "20-ott-2026"),
              ("20/10/2026", 15, "ott 2026"), ("20/10/2026 13:30", 20, "20/10/2026 13:30"),
              ("13:30", 16, "1:30 PM"), ("13:30:05", 19, "13:30:05"), ("abc", 4, "abc"), ("12", 0, "12"),
              ("1234,5", 24, "1.234,50 €")]
    for testo, codice, atteso in attesi:
        assert mostra(testo, codice) == atteso, (testo, codice, mostra(testo, codice), atteso)
    assert EditorExcel.lettera_colonna(0) == "A" and EditorExcel.lettera_colonna(27) == "AB"


class Clic:
    """Evento finto per i clic sulla tabella."""

    def __init__(self, widget, stato=0):
        self.widget, self.state = widget, stato


@prova
def finestra_colonne_righe_salva_anteprima_applica():
    import EditorExcel
    import EditorModelli
    from openpyxl import load_workbook
    from tkinter import filedialog, messagebox
    EditorModelli.FILE_PREFERENZE = os.path.join(cartella, "preferenze.json")
    lavoro = os.path.join(cartella, "lavoro")
    os.makedirs(lavoro)
    esempio = os.path.join(cartella, "listino.txt")
    with open(esempio, "w", encoding="cp1252") as f:
        f.write(TESTO)
    modello = os.path.join(cartella, "modelli", "listino.gui")
    os.makedirs(os.path.dirname(modello))
    risposte = {"askopenfilename": esempio, "asksaveasfilename": modello, "askopenfilenames": (esempio,)}
    for nome, valore in risposte.items():
        setattr(filedialog, nome, lambda *a, _v=valore, **k: _v)
    messaggi = []
    for nome in ("showinfo", "showerror"):
        setattr(messagebox, nome, lambda titolo, testo, _n=nome, **k: messaggi.append((_n, testo)))
    messagebox.askokcancel = lambda *a, **k: True

    app = EditorExcel.EditorExcel()
    app.config_servizio = {"cartella_lavoro": lavoro, "cartella_template": "", "cartella_output": lavoro}
    try:
        app.nuovo()
        app.update()
        assert len(app.colonne) == 6 and len(app.righe) == 5, (len(app.colonne), len(app.righe))
        # Modello nuovo: la riga 1 è l'intestazione, in grassetto e come testo
        assert [s["righe"] for s in app.stili] == ["1"] and app.stili[0]["testo"]
        assert app.celle[0][2].cget("text") == "Quantità" and "bold" in str(app.celle[0][2].cget("font"))

        # Colonne C e D: intero con migliaia, Courier grassetto rosso, a destra, larghe 12
        app.seleziona([2, 3])
        app.formato.set(EditorExcel.NOMI_FORMATI[3])
        app._formato_cambiato()
        app.gruppo_celle.font.set("Courier")
        app.gruppo_celle.grassetto.set(True)
        app.gruppo_celle.colore.set("rosso")
        app._carattere_cambiato("carattere", app.gruppo_celle)
        app.allineamento.set("destra")
        app._allineamento_cambiato()
        app.larghezza.set("12")
        app._larghezza_cambiata()
        # Solo la D: valuta con decimali
        app.seleziona([3])
        app.formato.set(EditorExcel.NOMI_FORMATI[8])
        app._formato_cambiato()
        assert "non si possono convertire" not in app.avviso.cget("text")
        # E: data; "non so" resta testo e l'editor lo dice (l'intestazione non conta: è come testo)
        app.seleziona([4])
        app.formato.set(EditorExcel.NOMI_FORMATI[12])
        app._formato_cambiato()
        assert "1 valori non si possono convertire" in app.avviso.cget("text"), app.avviso.cget("text")

        # Intestazione: Arial 12 grassetto blu su sfondo grigio
        app.lista_stili.selection_set(0)
        app._scegli_stile()
        app.gruppo_stile.font.set("Helvetica")
        app.gruppo_stile.dimensione.set("12")
        app.gruppo_stile.colore.set("blu")
        app._carattere_stile_cambiato()
        app.sfondo_stile.set("grigio chiaro")
        app._sfondo_stile_cambiato()

        # Totale sull'ultima riga: clic sul numero di riga 5, nuovo stile, poi "u1" (ultima riga)
        app._clic_numero_riga(Clic(app.numeri_riga[4]))
        assert app.righe_selezionate == {5}
        app.nuovo_stile()
        assert app.stili[1]["righe"] == "5" and app.stile_corrente == 1
        app.righe_stile.set("u1")
        app._righe_stile_cambiate()
        app.nome_stile.set("  Totale   generale ")
        app._nome_stile_cambiato()
        app.sfondo_stile.set("giallo chiaro")
        app._sfondo_stile_cambiato()
        # Righe scritte male: non cambiano lo stile
        app.righe_stile.set("u0")
        app._righe_stile_cambiate()
        assert app.stili[1]["righe"] == "u1" and "Righe non valide" in app.stato.cget("text")
        app.update()
        assert app.celle[2][3].cget("text") == "1.234,50 €" and app.celle[2][3].cget("anchor") == "e"
        # Nella riga del totale il carattere è quello dello stile (Calibri grassetto), il formato della colonna
        assert app.celle[4][3].cget("text") == "1.222,65 €" and app.celle[4][3].cget("background") == "#FFF2CC"
        assert app.celle[0][3].cget("background") == "#D9D9D9"
        # La riga 5 è ancora scelta (colore della selezione); tolta la scelta si vede che è speciale
        assert app.numeri_riga[4].cget("background") == EditorExcel.SFONDO_LETTERE_SELEZIONE
        app.righe_selezionate = set()
        app._aggiorna_numeri_riga()
        assert app.numeri_riga[4].cget("background") == EditorExcel.SFONDO_RIGA_SPECIALE

        # Annulla / ripeti (anche degli stili)
        prima = app._istantanea()
        app.annulla()
        assert app.stili[1]["sfondo"] is None
        app.ripeti()
        assert app._istantanea() == prima
        app.sposta_stile(-1)
        assert app.stili[0]["nome"] == "Totale generale" and app.stile_corrente == 0
        app.annulla()
        assert app.stili[0]["nome"] == "Intestazione"

        # Copia il formato della C sulla A, poi annulla
        app.seleziona([2])
        app.copia_colonna()
        app.seleziona([0])
        app.incolla_colonna()
        assert app.colonne[0] == app.colonne[2]
        app.annulla()
        assert app.colonne[0] == modello_gui.colonna_excel_vuota()

        # Salva e riapri
        assert app.salva(), messaggi
        intestazione, linee = modello_gui.leggi_gui(modello)
        assert intestazione["azione"] == "excel" and intestazione["uscita"] == "{nome}.xlsx", intestazione
        assert linee == ["1\tnome=Intestazione\tcarattere=hsn12b\tsfondo=grigio\ttesto=si",
                         "u1\tnome=Totale generale\tcarattere=ksn11n\tsfondo=giallo",
                         "colonna",
                         "0", "0",
                         "3\tcarattere=csn11r\tlarghezza=12\tallineamento=destra",
                         "8\tcarattere=csn11r\tlarghezza=12\tallineamento=destra",
                         "12", "fine"], linee
        stato = app._istantanea()
        app.apri_modello(modello)
        assert app._istantanea() == stato
        assert app.preferenze["recenti_excel"][0] == os.path.abspath(modello)

        # Anteprima con il codice del servizio
        xlsx = app.crea_documento(tempfile.mkdtemp(dir=cartella))
        ws = load_workbook(xlsx).worksheets[0]
        assert ws["D3"].value == 1234.5 and ws["D3"].number_format == modello_gui.FORMATI_EXCEL[8][1]
        assert (ws["D3"].font.name, ws["D3"].font.bold, ws["D3"].font.color.rgb) == ("Courier New", True, "FFFF0000")
        assert (ws["D1"].font.name, ws["D1"].font.size, ws["D1"].font.color.rgb) == ("Arial", 12, "FF0000FF")
        assert ws["D1"].value == "Prezzo" and ws["D1"].fill.start_color.rgb == "FFD9D9D9"
        assert (ws["D5"].value, ws["D5"].font.name, ws["D5"].fill.start_color.rgb) == (1222.65, "Calibri", "FFFFF2CC")
        assert ws["C2"].alignment.horizontal == "right" and ws.column_dimensions["C"].width == 12
        assert ws["E2"].value == datetime.datetime(2026, 10, 20) and ws["E4"].value == "non so"

        # Applica: il servizio elabora il .gui creato nella cartella di lavoro
        app.applica()
        assert os.path.isfile(os.path.join(lavoro, "listino.gui")), os.listdir(lavoro)
        sys.path.insert(0, os.path.join(CARTELLA_PROGRAMMA, "test"))
        import esegui_test
        servizio = esegui_test.carica_programma("SingleGUY")
        servizio.main("listino.txt", dict(app.config_servizio))
        ws = load_workbook(os.path.join(lavoro, "listino.xlsx")).worksheets[0]
        assert ws["D3"].value == 1234.5 and ws["D3"].font.name == "Courier New"
        assert ws["D5"].fill.start_color.rgb == "FFFFF2CC"

        # Un modello PDF non si apre nell'editor Excel
        pdf = os.path.join(cartella, "modelli", "fattura.gui")
        modello_gui.scrivi_gui(pdf, modello_gui.nuovo_modello(False), ["001\ttnn10", "fine"])
        app.apri_modello(pdf)
        assert "EditorModelli" in messaggi[-1][1] and app.percorso_modello == modello
    finally:
        app.modificato = False
        app.destroy()


@prova
def finestra_modalita_righe():
    """Con la scheda "Righe speciali" aperta il clic sceglie righe e apre il loro stile."""
    import EditorExcel
    import EditorModelli
    EditorModelli.FILE_PREFERENZE = os.path.join(cartella, "preferenze_righe.json")
    app = EditorExcel.EditorExcel()
    try:
        app._mostra_testo(TESTO)
        app.stili.append(dict(modello_gui.stile_riga_vuoto(), righe="u1", nome="Totale", sfondo="giallo"))
        app._aggiorna_tutto()
        app.update()
        # Scheda Colonne: il clic su una cella sceglie la colonna
        app._clic_cella(Clic(app.celle[2][3]))
        assert app.selezionate == [3] and not app.righe_selezionate
        app.schede.select(1)
        app.update()
        assert app._modalita_righe()
        # Ora lo stesso clic sceglie la riga 3, che non ha stile: nessuno stile aperto
        app._clic_cella(Clic(app.celle[2][3]))
        assert app.righe_selezionate == {3} and app.stile_corrente is None
        assert "Nuovo" in app.stato.cget("text")
        assert app.celle[2][0].cget("highlightbackground") == EditorExcel.BORDO_RIGA_SCELTA
        assert app.celle[2][3].cget("background") == EditorExcel.SFONDO, "la colonna non va evidenziata"
        # Clic sull'ultima riga: si apre il suo stile, che si modifica subito
        app._clic_cella(Clic(app.celle[4][1]))
        assert app.righe_selezionate == {5} and app.stili[app.stile_corrente]["nome"] == "Totale"
        assert app.nome_stile.get() == "Totale"
        app.sfondo_stile.set("azzurro chiaro" if "azzurro chiaro" in EditorExcel.NOMI_SFONDI else "azzurro")
        app._sfondo_stile_cambiato()
        assert app.stili[1]["sfondo"] == "azzurro"
        # Il bordo segna la riga scelta ma lo sfondo (nuovo) si vede
        assert app.celle[4][1].cget("background") == "#DDEBF7"
        assert app.celle[4][1].cget("highlightbackground") == EditorExcel.BORDO_RIGA_SCELTA
        # Ctrl+clic su una riga di un altro stile: stili diversi, nessuno aperto
        app._clic_cella(Clic(app.celle[0][0], stato=0x4))
        assert app.righe_selezionate == {1, 5} and app.stile_corrente is None
        # Maiusc+clic e "Nuovo": uno stile con le righe 2-4
        app._clic_cella(Clic(app.celle[1][0]))
        app._clic_cella(Clic(app.celle[3][0], stato=0x1))
        assert app.righe_selezionate == {2, 3, 4}
        app.nuovo_stile()
        assert app.stili[-1]["righe"] == "2-4" and app.stile_corrente == len(app.stili) - 1
        # Frecce su e giù
        app._sposta_riga(-1, Clic(None))
        assert app.righe_selezionate == {1}
        # Clic sulla lettera di una colonna: si torna alle colonne
        app._clic_cella(Clic(app.lettere[2]))
        app.update()
        assert not app._modalita_righe() and app.selezionate == [2]
        assert app.celle[0][0].cget("highlightbackground") != EditorExcel.BORDO_RIGA_SCELTA
    finally:
        app.modificato = False
        app.destroy()


shutil.rmtree(cartella, ignore_errors=True)
for nome, errore in risultati:
    print("[%s] %s" % ("OK " if errore is None else "ERR", nome))
    if errore:
        print("       " + errore.replace("\n", "\n       "))
print("\n%d/%d prove OK" % (sum(e is None for _, e in risultati), len(risultati)))
sys.exit(0 if all(e is None for _, e in risultati) else 1)
