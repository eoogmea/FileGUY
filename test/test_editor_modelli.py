"""Prova dell'editor dei modelli (EditorModelli.py / modello_gui.py).

Uso:  python test\\test_editor_modelli.py

1. Formati -> linee del .gui -> formati: deve tornare uguale.
2. Il servizio (SingleGUY) applica davvero i formati scelti, nel PDF e nel Word.
3. La finestra: formattazione di una selezione, salvataggio, anteprima e "Applica a file di testo".
"""
import os
import random
import shutil
import sys
import tempfile
import traceback

import pymupdf
from docx import Document

CARTELLA_PROGRAMMA = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, CARTELLA_PROGRAMMA)
os.chdir(CARTELLA_PROGRAMMA)
import modello_gui

risultati = []
cartella = tempfile.mkdtemp(prefix="infas_modelli_")
CARTELLA_TEMPLATE = os.path.join(cartella, "template")
os.makedirs(CARTELLA_TEMPLATE)
modello_pdf = pymupdf.open()
modello_pdf.new_page(width=595, height=842).draw_rect(pymupdf.Rect(20, 20, 575, 822), color=(0, 0, 1))
modello_pdf.save(os.path.join(CARTELLA_TEMPLATE, "modello.pdf"))

RIGHE = [
    "INFAS SERVICE S.r.l.",
    "",
    "Fattura n. 123 del 20/10/2026",
    "Cliente: Mario Rossi   Totale: 1.234,56 €",
    "Descrizione con accenti: città, perché, più, così",
    "Testo bianco su fondo scuro",
    "Verdana: e Segoe UI",
]


def prova(funzione):
    try:
        funzione()
        risultati.append((funzione.__name__, None))
    except Exception:
        risultati.append((funzione.__name__, traceback.format_exc(limit=-4).strip()))
    return funzione


def formati_di_prova(word):
    t = ("t", True, False, 14, "n")
    h = ("h", False, True, 10, "r")
    c = ("c", True, True, 12, "b")
    k = ("k", True, False, 11, "v")
    g = ("g", False, False, 10, "w")
    v = ("v", False, True, 9, "n")
    u = ("u", True, True, 10, "n")
    return [
        [t] * len(RIGHE[0]),
        [],
        [c] * len(RIGHE[2]),
        [h] * 8 + [t] * 13 + [c] * (len(RIGHE[3]) - 21),
        [k] * len(RIGHE[4]),
        [g] * len(RIGHE[5]),
        [v] * 9 + [u] * (len(RIGHE[6]) - 9),
    ]


@prova
def andata_e_ritorno_casuale():
    casuale = random.Random(1)
    for word in (False, True):
        for _ in range(200):
            righe, formati = [], []
            for _ in range(casuale.randint(1, 8)):
                lunghezza = casuale.randint(0, 99)
                righe.append("x" * lunghezza)
                riga, formato = [], None
                for _ in range(lunghezza):
                    if formato is None or casuale.random() < 0.1:
                        formato = (casuale.choice("thckvgu"), casuale.random() < .5, casuale.random() < .5,
                                   casuale.randint(6, 36), casuale.choice("nnnrvbw"))
                    riga.append(formato)
                formati.append(riga)
            linee = modello_gui.linee_formato(formati, word)
            assert modello_gui.leggi_formati(linee, righe, word) == formati, linee


@prova
def colonna_oltre_99():
    riga = [modello_gui.PREDEFINITO] * 100 + [("h", True, False, 10, "n")] * 5
    try:
        modello_gui.linee_formato([riga], False)
        raise AssertionError("un formato dopo la colonna 99 è stato accettato")
    except modello_gui.ErroreModello:
        pass
    # Un formato solo, anche su una riga lunga, va bene
    assert modello_gui.linee_formato([[modello_gui.PREDEFINITO] * 150], False) == ["001\ttnn10", "fine"]


def servizio():
    import importlib
    return importlib.import_module("SingleGUY")


def crea_con_servizio(word, nome):
    intestazione = modello_gui.nuovo_modello(word)
    intestazione.update(testo=nome + ".txt", uscita=nome + (".docx" if word else ".pdf"),
                        template="" if word else "modello.pdf")
    lavoro = os.path.join(cartella, nome)
    os.makedirs(lavoro)
    with open(os.path.join(lavoro, nome + ".txt"), "w") as f:
        f.write("\n".join(RIGHE) + "\n")
    modello_gui.scrivi_gui(os.path.join(lavoro, nome + ".gui"), intestazione,
                           modello_gui.linee_formato(formati_di_prova(word), word))
    with open(os.path.join(lavoro, nome + ".gui")) as f:
        guida = f.readlines()
    config = {"cartella_lavoro": lavoro, "cartella_output": lavoro, "cartella_template": CARTELLA_TEMPLATE}
    if word:
        servizio().docx(guida, config)
    else:
        servizio().pdf(guida, config)
    return os.path.join(lavoro, intestazione["uscita"])


def famiglia_nel_pdf(nome_pdf):
    """La famiglia (senza spazi) del font che il servizio usa davvero per un nome del PDF: dipende dal
    computer, per esempio su Linux al posto di Calibri c'è Helvetica e al posto di Consolas DejaVu Sans Mono."""
    import re
    nome = pymupdf.Font(fontfile=servizio().trova_font(nome_pdf)).name
    # Lo stile dopo la famiglia: "Regular", "Bold", ma DejaVu chiama "Book" il normale
    return re.sub(r"\s+(Regular|Book|Bold|Italic|Oblique)\b.*$", "", nome).replace(" ", "")


@prova
def email_di_prima_nel_formato_nuovo():
    def intestazione(destinatario, coda):
        return dict(modello_gui.nuovo_modello(False), invio="email", destinatario=destinatario, coda=coda)

    vecchia = intestazione("", ["c@x.it", "Oggetto {nome}", "", "Riga 1", "Riga 2"])
    assert modello_gui.aggiorna_email(vecchia)
    assert vecchia["destinatario"] == "c@x.it"
    assert vecchia["coda"] == ["Oggetto {nome}", "", "Riga 1", "Riga 2"], vecchia["coda"]
    senza_testo = intestazione("", ["c@x.it", "Oggetto"])
    assert modello_gui.aggiorna_email(senza_testo) and senza_testo["coda"] == ["Oggetto"]
    con_allegato = intestazione("", ["c@x.it", "Oggetto", "C:\\allegato.pdf", "Testo"])
    assert modello_gui.aggiorna_email(con_allegato)
    assert con_allegato["coda"] == ["Oggetto", "allegato: C:\\allegato.pdf", "", "Testo"], con_allegato["coda"]
    # Già nuova o senza email: resta com'è
    for invariata in (intestazione("c@x.it", ["Oggetto", "", "Testo"]),
                      dict(intestazione("", ["doc.pdf", "Testo"]), invio="telegram")):
        prima = dict(invariata)
        assert not modello_gui.aggiorna_email(invariata) and invariata == prima


@prova
def servizio_applica_formati_pdf():
    doc = pymupdf.open(crea_con_servizio(False, "pdf"))
    pezzi = {}
    for blocco in doc[0].get_text("dict")["blocks"]:
        for riga in blocco.get("lines", []):
            for span in riga["spans"]:
                pezzi[span["text"].strip()] = (span["font"], round(span["size"]), "%06X" % span["color"])
    # testo: (font del .gui come lo chiama il servizio, grassetto, dimensione, colore)
    attesi = {
        "INFAS SERVICE S.r.l.": ("Times", True, 14, "000000"),
        "Fattura n. 123 del 20/10/2026": ("Courier", True, 12, "0000FF"),
        "Cliente:": ("Helvetica", False, 10, "FF0000"),
        "Mario Rossi": ("Times", True, 14, "000000"),
        "Totale: 1.234,56 €": ("Courier", True, 12, "0000FF"),
        "Descrizione con accenti: città, perché, più, così": ("Calibri", True, 11, "008000"),
        "Testo bianco su fondo scuro": ("Georgia", False, 10, "FFFFFF"),
        "Verdana:": ("Verdana", False, 9, "000000"),
        "e Segoe UI": ("SegoeUI", True, 10, "000000"),
    }
    for testo, (nome_pdf, grassetto, dimensione, colore) in attesi.items():
        font = famiglia_nel_pdf(nome_pdf)
        assert testo in pezzi, "manca %r in %r" % (testo, list(pezzi))
        trovato_font, trovata_dimensione, trovato_colore = pezzi[testo]
        assert trovato_font.replace(" ", "").startswith(font) and trovata_dimensione == dimensione \
            and trovato_colore == colore and ("Bold" in trovato_font) == grassetto, \
            "%r: %s %d %s invece di %s %d %s" % (testo, trovato_font, trovata_dimensione, trovato_colore, font,
                                               dimensione, colore)
    # Nel PDF ci sono solo i font usati
    usati = {f[3] for f in doc.get_page_fonts(0)}
    assert not any("Courier" in f and "Italic" not in f for f in usati), usati


@prova
def servizio_applica_formati_word():
    doc = Document(crea_con_servizio(True, "word"))
    runs = [(r.text, r.font.name, r.font.size.pt if r.font.size else None, r.font.bold, r.font.italic,
             str(r.font.color.rgb) if r.font.color.type else None) for p in doc.paragraphs for r in p.runs]
    testo = "".join(r[0] for r in runs)
    assert testo == "\n".join(RIGHE) + "\n", repr(testo)
    def run_con(inizio):
        return next(r for r in runs if r[0].strip().startswith(inizio.strip()))
    assert run_con("INFAS")[1:] == ("Times New Roman", 14, True, False, "000000"), run_con("INFAS")
    assert run_con("Fattura")[1:] == ("Courier New", 12, True, True, "0000FF"), run_con("Fattura")
    assert run_con("Cliente:")[1:] == ("Arial", 10, False, True, "FF0000"), run_con("Cliente:")
    assert run_con(" Mario Rossi")[1:] == ("Times New Roman", 14, True, False, "000000"), run_con(" Mario")
    assert run_con("Totale")[1:] == ("Courier New", 12, True, True, "0000FF"), run_con("Totale")
    assert run_con("Descrizione")[1:] == ("Calibri", 11, True, False, "008000"), run_con("Descrizione")
    assert run_con("Testo bianco")[1:] == ("Georgia", 10, False, False, "FFFFFF"), run_con("Testo bianco")
    assert run_con("Verdana:")[1:] == ("Verdana", 9, False, True, "000000"), run_con("Verdana:")
    assert run_con("e Segoe UI")[1:] == ("Segoe UI", 10, True, True, "000000"), run_con("e Segoe UI")


@prova
def finestra_filigrana():
    """La filigrana segue le modifiche, evidenzia la riga del cursore e un clic porta il cursore lì."""
    import EditorModelli
    from tkinter import filedialog
    esempio = os.path.join(cartella, "esempio_filigrana.txt")
    # Abbastanza righe per avere due pagine (44 righe per pagina A4)
    righe = RIGHE * 8
    with open(esempio, "w") as f:
        f.write("\n".join(righe) + "\n")
    filedialog.askopenfilename = lambda *a, **k: esempio
    app = EditorModelli.EditorModelli()
    app.config_servizio = {"cartella_lavoro": cartella, "cartella_template": CARTELLA_TEMPLATE,
                           "cartella_output": cartella}
    try:
        app.nuovo()
        app.template.set("modello.pdf")
        app.testo.tag_add("sel", "4.0", "4.end")
        app._imposta("f", "c")
        app.testo.tag_remove("sel", "1.0", "end")
        app.filigrana()
        finestra = app.finestra_filigrana
        assert finestra is not None and finestra.documento is not None, finestra.stato.cget("text")
        assert finestra.documento.page_count == 2 and finestra.righe_per_pagina == 44
        finestra.update()
        scala = finestra._scala()
        # Clic a metà del carattere della colonna 9 di "Cliente: Mario Rossi" (riga 4) -> riga 3, colonna 9
        x0, x1 = finestra.caratteri[0][3][9]
        y = finestra._y_riga(3) + 8
        evento = type("Evento", (), {"x": EditorModelli.MARGINE_IMMAGINE + (x0 + x1) / 2 * scala + 1,
                                     "y": EditorModelli.MARGINE_IMMAGINE + y * scala})()
        finestra.tela.yview_moveto(0)
        finestra.tela.xview_moveto(0)
        finestra._clic(evento)
        assert app.testo.index("insert") == "4.10", app.testo.index("insert")
        assert len(finestra.caratteri[0][3]) == len(RIGHE[3])
        # Il cursore su una riga della seconda pagina porta la filigrana lì
        app.vai_a(50, 0)
        assert finestra.pagina == 1 and finestra.tela.find_withtag("cursore")
        # Una modifica ricrea la filigrana dopo un attimo
        documento = finestra.documento
        app.testo.tag_add("sel", "1.0", "1.end")
        app._alterna("b")
        app.after(900, app.quit)
        app.mainloop()
        assert finestra.documento is not documento
        # Template orizzontale: 29 righe per pagina e le posizioni restano quelle delle righe giuste
        orizzontale = pymupdf.open()
        orizzontale.new_page(width=842, height=595).insert_text((300, 300), "Testo del template")
        orizzontale.save(os.path.join(CARTELLA_TEMPLATE, "orizzontale.pdf"))
        app.template.set("orizzontale.pdf")
        finestra.aggiorna()
        assert finestra.righe_per_pagina == 29 and finestra.documento.page_count == 2, finestra.stato.cget("text")
        # La riga 31 (indice 30) è la seconda della seconda pagina
        assert righe[30] and len(finestra.caratteri[1][30]) == len(righe[30])
        finestra._chiudi()
        assert app.finestra_filigrana is None
    finally:
        app.modificato = False
        app.destroy()


@prova
def finestra_formatta_salva_anteprima_applica():
    import EditorModelli
    from tkinter import filedialog, messagebox
    lavoro = os.path.join(cartella, "lavoro_finestra")
    os.makedirs(lavoro)
    esempio = os.path.join(cartella, "esempio.txt")
    with open(esempio, "w") as f:
        f.write("\n".join(RIGHE) + "\n")
    percorso_modello = os.path.join(cartella, "modelli", "fattura.gui")
    os.makedirs(os.path.dirname(percorso_modello))
    risposte = {"askopenfilename": esempio, "asksaveasfilename": percorso_modello,
                "askopenfilenames": (esempio,)}
    for nome, valore in risposte.items():
        setattr(filedialog, nome, lambda *a, _v=valore, **k: _v)
    messaggi = []
    for nome in ("showinfo", "showerror"):
        setattr(messagebox, nome, lambda titolo, testo, _n=nome, **k: messaggi.append((_n, testo)))
    messagebox.askokcancel = lambda *a, **k: True

    app = EditorModelli.EditorModelli()
    app.config_servizio = {"cartella_lavoro": lavoro, "cartella_template": CARTELLA_TEMPLATE,
                           "cartella_output": lavoro}
    try:
        app.nuovo()
        assert app.righe[:len(RIGHE)] == RIGHE
        # Riga 1 intera: Courier 12 grassetto; riga 4 colonne 10-21 ("Mario Rossi"): Helvetica corsivo
        app.testo.tag_add("sel", "1.0", "1.end")
        app.scelta_font.set("Courier")
        app._imposta("f", "c")
        app.scelta_dimensione.set("12")
        app._imposta_dimensione()
        app._alterna("b")
        app.testo.tag_remove("sel", "1.0", "end")
        app.testo.tag_add("sel", "4.9", "4.20")
        app._imposta("f", "h")
        app._alterna("i")
        # Solo alcune lettere ("Mario", colonne 9-14) in rosso
        app.testo.tag_remove("sel", "1.0", "end")
        app.testo.tag_add("sel", "4.9", "4.14")
        app._imposta("c", "r")
        app.update()
        assert app.modificato
        app.template.set("modello.pdf")
        assert app.salva(), messaggi
        assert os.path.isfile(percorso_modello) and os.path.isfile(percorso_modello[:-4] + ".txt")
        intestazione, linee = modello_gui.leggi_gui(percorso_modello)
        nome_testo = os.path.splitext(os.path.basename(percorso_modello))[0] + ".txt"
        assert intestazione["testo"] == nome_testo and intestazione["uscita"] == "{nome}.pdf", intestazione
        assert linee[0] == "001\tcsn12", linee
        # Riga con un colore: codici di 6 caratteri; righe senza colori: 5 come prima
        assert linee[3] == "004\ttnn10n 00 09\thns10r 09 14\thns10n 14 20\ttnn10n 20 41", linee
        assert linee[4] == "005\ttnn10", linee

        # Riaprendo il modello la formattazione torna uguale
        formati = app._formati_testo()
        app.apri_modello(percorso_modello)
        assert app._formati_testo() == formati

        # Anteprima PDF con il codice del servizio
        pdf = app.crea_documento(tempfile.mkdtemp(dir=cartella))
        assert pdf and pymupdf.open(pdf)[0].get_text().count("Mario Rossi") == 1
        app.FinestraAnteprima = EditorModelli.FinestraAnteprima(app, pdf)
        app.update()
        app.FinestraAnteprima._chiudi()

        # Anteprima Word
        app.word.set(True)
        app._cambia_tipo()
        assert app.uscita.get() == "{nome}.docx" and app.azione.get() == "docx"
        docx = app.crea_documento(tempfile.mkdtemp(dir=cartella))
        assert docx and "Mario Rossi" in "".join(p.text for p in Document(docx).paragraphs)
        app.word.set(False)
        app._cambia_tipo()

        # Applica: crea esempio.gui nella cartella di lavoro con {nome} sostituito
        app.salva()
        app.applica()
        gui = os.path.join(lavoro, "esempio.gui")
        assert os.path.isfile(gui) and os.path.isfile(os.path.join(lavoro, "esempio.txt")), os.listdir(lavoro)
        intestazione, _ = modello_gui.leggi_gui(gui)
        assert intestazione["testo"] == "esempio.txt" and intestazione["uscita"] == "esempio.pdf", intestazione

        # Il servizio elabora il .gui creato
        modulo = servizio()
        modulo.main("esempio.txt", dict(app.config_servizio))
        assert os.path.isfile(os.path.join(lavoro, "esempio.pdf")), os.listdir(lavoro)
    finally:
        app.modificato = False
        app.destroy()


@prova
def finestra_aiuti_alla_formattazione():
    """Annulla/ripeti, copia formato, blocco, formato di una riga su altre, righello, zoom, recenti."""
    import EditorModelli
    EditorModelli.FILE_PREFERENZE = os.path.join(cartella, "preferenze", "editor_modelli.json")
    righe = ["Codice  Descrizione          Q.ta   Importo",
             "A001    Vite                 10     12,50",
             "A002    Bullone lungo         5      7,00",
             "A003    Rondella            100      1,20"]
    app = EditorModelli.EditorModelli()
    try:
        app._mostra_testo(righe)
        app.update()

        def linee():
            return modello_gui.linee_formato(app._formati_testo(), False)
        iniziali = linee()

        # Annulla e ripeti della formattazione
        app.testo.tag_add("sel", "1.0", "1.end")
        app._alterna("b")
        con_grassetto = linee()
        assert con_grassetto[0] == "001\ttsn10", con_grassetto
        app.annulla_formato()
        assert linee() == iniziali
        app.ripeti_formato()
        assert linee() == con_grassetto
        # Ctrl+Z senza "Modifica testo" annulla il formato, non il testo
        app._annulla_o_ripeti(False)
        assert linee() == iniziali and app.righe == righe

        # Selezione a blocco: colonne 37-42 (gli importi) delle righe 2-4 in Courier;
        # le righe degli articoli finiscono alla colonna 41: il blocco si ferma lì
        app.seleziona_blocco(2, 4, 37, 42)
        assert len(app.testo.tag_ranges("sel")) == 6
        app._imposta("f", "c")
        risultato = linee()
        for n in (1, 2, 3):
            assert risultato[n] == "%03d\ttnn10 00 37\tcnn10 37 41" % (n + 1), risultato
        assert risultato[0] == "001\ttnn10", risultato

        # Copia e incolla formato: il formato degli importi sulla riga 1 colonne 0-6
        app.testo.tag_remove("sel", "1.0", "end")
        app.testo.mark_set("insert", "2.40")
        app.copia_formato()
        assert app.formato_copiato == ("c", False, False, 10, "n"), app.formato_copiato
        app.testo.tag_add("sel", "1.0", "1.6")
        app.incolla_formato()
        assert linee()[0] == "001\tcnn10 00 06\ttnn10 06 %d" % len(righe[0]), linee()[0]

        # Formato della riga 2 sulle righe 3-4 (già uguali) e sulla riga 1
        app.testo.tag_add("sel", "2.0", "2.4")
        app._alterna("b")
        app.copia_formato_riga(2, 1, 4)
        # L'ultimo formato arriva a fine riga, anche dove la riga è più lunga della riga 2
        risultato = linee()
        assert risultato[0] == "001\ttsn10 00 04\ttnn10 04 37\tcnn10 37 %d" % len(righe[0]), risultato
        assert risultato[2] == "003\ttsn10 00 04\ttnn10 04 37\tcnn10 37 41", risultato

        # Il righello disegna le colonne della riga del cursore
        app.testo.mark_set("insert", "2.0")
        app._disegna_righello()
        app.update()
        assert app.righello.find_all(), "righello vuoto"

        # Lo zoom cambia solo la vista: il .gui resta uguale
        prima = linee()
        app.cambia_zoom(1)
        app.cambia_zoom(1)
        assert app.zoom == 1.25 and linee() == prima
        nomi_vista = [t for t in app.testo.tag_names() if t.startswith("v|")]
        assert any("|12|" in t for t in nomi_vista), nomi_vista
        app.cambia_zoom(0)
        assert app.zoom == 1.0

        # Modelli recenti: salvati e riproposti, il più recente per primo
        modello = os.path.join(cartella, "recenti", "articoli.gui")
        os.makedirs(os.path.dirname(modello))
        from tkinter import filedialog
        filedialog.asksaveasfilename = lambda *a, **k: modello
        app.percorso_testo = os.path.join(cartella, "articoli_esempio.txt")
        with open(app.percorso_testo, "w") as f:
            f.write("\n".join(righe))
        assert app.salva(con_nome=True)
        app._elenca_recenti()
        assert app.menu_recenti.entrycget(0, "label").endswith("articoli.gui")
        nuova = EditorModelli.EditorModelli()
        try:
            assert nuova.preferenze["recenti"][0] == os.path.abspath(modello)
        finally:
            nuova.destroy()
        formati = app._formati_testo()
        app.apri_recente(modello)
        assert app._formati_testo() == formati
    finally:
        app.modificato = False
        app.destroy()


@prova
def font_a_larghezza_fissa_nel_pdf_e_in_word():
    """Ogni font a larghezza fissa finisce davvero nel PDF (e con il suo nome in Word)."""
    import re
    righe = ["Riga %s grassetto corsivo 0123456789" % lettera for lettera in modello_gui.FONT_FISSI]
    formati = [[(lettera, True, True, 11, "n")] * len(righe[n]) for n, lettera in enumerate(modello_gui.FONT_FISSI)]
    cartella_prova = tempfile.mkdtemp(dir=cartella)
    with open(os.path.join(cartella_prova, "fissi.txt"), "w") as f:
        f.write("\n".join(righe) + "\n")
    modulo = servizio()
    config = {"cartella_lavoro": cartella_prova, "cartella_template": CARTELLA_TEMPLATE,
              "cartella_output": cartella_prova}
    for word, uscita in ((False, "fissi.pdf"), (True, "fissi.docx")):
        intestazione = dict(modello_gui.nuovo_modello(word), testo="fissi.txt", uscita=uscita)
        modello_gui.scrivi_gui(os.path.join(cartella_prova, "fissi.gui"), intestazione,
                               modello_gui.linee_formato(formati, word))
        with open(os.path.join(cartella_prova, "fissi.gui")) as f:
            guida = f.readlines()
        (modulo.docx if word else modulo.pdf)(guida, config)
    font_pdf = {re.sub(r"^[A-Z]{6}\+", "", s["font"]) for b in pymupdf.open(os.path.join(cartella_prova, "fissi.pdf"))[0]
                .get_text("dict")["blocks"] for l in b["lines"] for s in l["spans"]}
    # Il font che il servizio usa su questo computer (su Linux Consolas diventa DejaVu Sans Mono)
    for lettera in modello_gui.FONT_FISSI:
        atteso = famiglia_nel_pdf(modulo.FONT_CODICI[lettera][0])
        assert any(atteso in nome.replace(" ", "").replace("-", "") for nome in font_pdf), (atteso, font_pdf)
    nomi_word = {run.font.name for run in Document(os.path.join(cartella_prova, "fissi.docx")).paragraphs[0].runs}
    assert {"Courier New", "Consolas", "DejaVu Sans Mono", "IBM Plex Mono", "Liberation Mono"} <= nomi_word, nomi_word


@prova
def testo_segnaposto():
    import re
    righe = ["DITTA ESEMPIO S.r.l.                FATTURA N. 123/2026",
             "A001      Viti 4x30       100      0,15       15,00",
             "--------------------------------------------------", "", "Spett.le"]
    x = modello_gui.testo_segnaposto(righe, 3)
    assert x[0] == "xxxxxxxxxxxxxxx.x.x.xxx             xxxxxxxxx.xxxx/xxxxxxx", x[0]
    assert x[1] == "xxxxxxx   xxxxxxxxxxxx    xxxxxx   x,xxxxx    xx,xxxxx", x[1]
    assert x[2] == righe[2] and x[3] == "" and x[4] == "xxxxx.xxxxx"
    assert modello_gui.testo_segnaposto(["a-b  c"], 0, simboli=False) == ["xxx  x"]
    # In ogni caso i campi restano nella loro colonna, con almeno uno spazio prima
    for in_piu in (0, 5, 30):
        for simboli in (True, False):
            for prima, dopo in zip(righe, modello_gui.testo_segnaposto(righe, in_piu, simboli)):
                for campo in re.finditer(r"\S+(?: \S+)*", prima):
                    s = campo.start()
                    assert dopo[s] != " " and (s == 0 or dopo[s - 1] == " "), (in_piu, prima, dopo)


@prova
def finestra_segnaposto():
    import EditorModelli
    from tkinter import filedialog
    EditorModelli.FILE_PREFERENZE = os.path.join(cartella, "preferenze_x.json")
    esempio = os.path.join(cartella, "fattura_x.txt")
    with open(esempio, "w") as f:
        f.write("\n".join(RIGHE) + "\n")
    modello = os.path.join(cartella, "modelli_x", "fattura_x.gui")
    os.makedirs(os.path.dirname(modello))
    filedialog.askopenfilename = lambda *a, **k: esempio
    filedialog.asksaveasfilename = lambda *a, **k: modello
    app = EditorModelli.EditorModelli()
    try:
        assert app.nuovo()
        # Riga 4 in Courier grassetto, poi le x: la formattazione resta sulle stesse colonne
        app.testo.tag_add("sel", "4.0", "4.end")
        app._imposta("f", "c")
        app._alterna("b")
        prima = app._formati_testo()[3]
        assert app.segnaposto((2, True))
        assert app.righe[3] == modello_gui.testo_segnaposto([RIGHE[3]], 2)[0], app.righe[3]
        # Niente più lettere o cifre: solo x, spazi e simboli (: . , €)
        assert not any(c.isalnum() and c != "x" for c in app.righe[3]), app.righe[3]
        assert app._formati_testo()[3][:len(prima)] == prima
        assert app.testo_cambiato and app.modificato
        assert app.salva()
        with open(modello[:-4] + ".txt") as f:
            assert f.read().split("\n")[3] == app.righe[3]
        with open(esempio) as f:
            assert f.read().split("\n") == RIGHE + [""], "il file di partenza è cambiato"
        # Riaprendo il modello c'è il testo con le x
        app.apri_modello(modello)
        assert app.righe[3] == modello_gui.testo_segnaposto([RIGHE[3]], 2)[0]
    finally:
        app.modificato = False
        app.destroy()


shutil.rmtree(cartella, ignore_errors=True)
for nome, errore in risultati:
    print("[%s] %s" % ("OK " if errore is None else "ERR", nome))
    if errore:
        print("       " + errore.replace("\n", "\n       "))
print("\n%d/%d prove OK" % (sum(e is None for _, e in risultati), len(risultati)))
sys.exit(0 if all(e is None for _, e in risultati) else 1)
