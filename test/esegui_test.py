"""Debug completo di SingleGUY con file .gui di prova.

Uso:  python test\\esegui_test.py [SingleGUY] [cartella dei modelli PDF]

Con la cartella dei modelli (es. C:\\prova\\modelli) prova anche ogni template vero.

Per ogni caso crea in test\\_run\\<caso> le cartelle lavoro / template / output,
scrive il .txt e il .gui, poi esegue start() del programma come farebbe il servizio.
Stampa, email, Telegram, WhatsApp, editor Word, apertura dei file ed exe sono simulati:
non parte niente verso l'esterno, ma viene registrato cosa sarebbe stato inviato.
I .gui generati restano in test\\casi per poterli riusare a mano.
"""
import contextlib
import datetime
import email
import io
import os
import shutil
import sys
import traceback
import types

import pymupdf
from PIL import Image

CARTELLA_TEST = os.path.dirname(os.path.abspath(__file__))
CARTELLA_PROGRAMMA = os.path.dirname(CARTELLA_TEST)
CARTELLA_RUN = os.path.join(CARTELLA_TEST, "_run")
CARTELLA_CASI = os.path.join(CARTELLA_TEST, "casi")
sys.path.insert(0, CARTELLA_PROGRAMMA)

TESTO_PDF = [
    "INFAS SERVICE S.r.l.",
    "Via Roma 1 - 00100 Roma",
    "",
    "Fattura n. 123 del 20/10/2026",
    "Cliente: Mario Rossi   Totale: 1.234,56 €",
    "Descrizione con accenti: città, perché, più, così",
    "Riga normale senza formato",
]

# Formati PDF: <riga>\t<font t/h/c><grassetto s/n><corsivo s/n><dimensione 2 cifre>
# più formati sulla stessa riga: <codice> <da 2 cifre> <a 2 cifre>, separati da un tab
# (es. "tnn11 00 08" = caratteri da 0 a 8 in Times normale 11)
FORMATI_PDF = [
    "001\ttsn14",
    "002\thnn10",
    "003",
    "004\tcss12",
    "005\ttnn11 00 08\ttsn11 10 23",
    "006\thns10",
    "007",
    "fine",
]

# Formati Word: come il PDF più il colore (n nero, r rosso, b blu)
FORMATI_DOCX = [
    "001\ttsn14n",
    "002\thnn10b",
    "003",
    "004\tcss12r",
    "005\ttnn11n 00 08\ttsn11r 10 23",
    "006\thns10n",
    "007",
    "fine",
]


def gui(testo, azione, uscita, template="", rotazione="0", invio="0", destinatario="",
        formati=None, coda=None):
    righe = [testo, azione, uscita, template, rotazione, invio, destinatario]
    righe += formati if formati is not None else FORMATI_PDF
    righe += coda or []
    return righe


def casi(modelli=None):
    """Ogni caso: nome, file da mettere nella cartella di lavoro, controlli sul risultato."""
    c = []
    testo = "\n".join(TESTO_PDF) + "\n"

    def caso(nome, gui_righe, txt=testo, nome_gui=None, controllo=None, uscite=(), chiamate=(),
             rimossi=True, lavoro_extra=(), modello=None, errore_atteso=False):
        c.append(dict(nome=nome, gui=gui_righe, txt=txt, nome_gui=nome_gui or nome + ".gui",
                      controllo=controllo, uscite=list(uscite), chiamate=list(chiamate), rimossi=rimossi,
                      lavoro_extra=list(lavoro_extra), modello=modello, errore_atteso=errore_atteso))

    def testo_pdf(nome, *attesi, pagine=None):
        """Controlla che ogni testo atteso compaia una sola volta e il numero di pagine."""
        def controllo(config):
            doc = pymupdf.open(os.path.join(config["cartella_output"], nome))
            testo_doc = "".join(p.get_text() for p in doc)
            errori = []
            if pagine is not None and doc.page_count != pagine:
                errori.append("%s: %d pagine invece di %d" % (nome, doc.page_count, pagine))
            for n, p in enumerate(doc):
                if p.get_text().strip() == "":
                    errori.append("%s: pagina %d vuota" % (nome, n + 1))
            doc.close()
            for a in attesi:
                volte = testo_doc.count(a)
                if volte != 1:
                    errori.append("%s: %r compare %d volte" % (nome, a, volte))
            return errori
        return controllo

    controllo_testo = testo_pdf("pdf_template.pdf", "Cliente:", "Mario Rossi", "1.234,56 €",
                                "città, perché, più, così", "Riga normale senza formato")

    # --- PDF ---
    caso("pdf_senza_template", gui("pdf_senza_template.txt", "pdf", "pdf_senza_template.pdf"),
         uscite=["pdf_senza_template.pdf"], chiamate=["startfile"],
         controllo=testo_pdf("pdf_senza_template.pdf", "Cliente:", "Mario Rossi", "1.234,56 €"))
    caso("pdf_template", gui("pdf_template.txt", "pdf", "pdf_template.pdf", "modello.pdf"),
         uscite=["pdf_template.pdf"], chiamate=["startfile"], controllo=controllo_testo)
    caso("pdf_template_ruotato", gui("pdf_template_ruotato.txt", "pdf", "pdf_template_ruotato.pdf",
                                     "modello.pdf", "90"),
         uscite=["pdf_template_ruotato.pdf"], chiamate=["startfile"])
    caso("pdf_senza_formati", gui("pdf_senza_formati.txt", "pdf", "pdf_senza_formati.pdf",
                                  formati=["001", "fine"]),
         uscite=["pdf_senza_formati.pdf"], chiamate=["startfile"])
    testo_lungo = "".join("Riga numero %d del documento lungo\n" % n for n in range(1, 121))
    caso("pdf_piu_pagine", gui("pdf_piu_pagine.txt", "file", "pdf_piu_pagine.pdf", "modello.pdf"),
         txt=testo_lungo, uscite=["pdf_piu_pagine.pdf"],
         controllo=testo_pdf("pdf_piu_pagine.pdf", "numero 1 ", "numero 61 ", "numero 120 ", pagine=3))
    caso("pdf_piu_pagine_senza_template", gui("pdf_piu_pagine_senza_template.txt", "file",
                                              "pdf_piu_pagine_senza_template.pdf"),
         txt=testo_lungo, uscite=["pdf_piu_pagine_senza_template.pdf"],
         controllo=testo_pdf("pdf_piu_pagine_senza_template.pdf", "numero 1 ", "numero 61 ", "numero 120 ",
                             pagine=3))
    caso("file", gui("file.txt", "file", "file.pdf", "modello.pdf"), uscite=["file.pdf"])
    caso("stampa", gui("stampa.txt", "stampa", "stampa.pdf", "modello.pdf"),
         uscite=["stampa.pdf"], chiamate=["stampa_pdf"])

    # --- Invii dopo il PDF ---
    caso("pdf_email", gui("pdf_email.txt", "file", "pdf_email.pdf", "modello.pdf", invio="email",
                          coda=["cliente@example.com", "Invio fattura 123", "",
                                "Buongiorno,", "in allegato la fattura."]),
         uscite=["pdf_email.pdf"], chiamate=["smtp_login", "smtp_sendmail"],
         controllo=lambda config: controlla_email("Invio fattura 123", ["pdf_email.pdf"],
                                                  "Buongiorno,\nin allegato la fattura.", "cliente@example.com"))
    caso("pdf_email_destinatario", gui("pdf_email_destinatario.txt", "file", "pdf_email_destinatario.pdf",
                                       "modello.pdf", invio="email", destinatario="cliente@example.com",
                                       coda=["Fattura 123", "", "Buongiorno,", "", "in allegato la fattura.",
                                             "Saluti"]),
         uscite=["pdf_email_destinatario.pdf"], chiamate=["smtp_login", "smtp_sendmail"],
         controllo=lambda config: controlla_email("Fattura 123", ["pdf_email_destinatario.pdf"],
                                                  "Buongiorno,\n\nin allegato la fattura.\nSaluti",
                                                  "cliente@example.com"))
    caso("pdf_email_2_allegati", gui("pdf_email_2_allegati.txt", "file", "pdf_email_2_allegati.pdf",
                                     invio="email",
                                     coda=["cliente@example.com", "Due allegati", "@ALLEGATO2@",
                                           "Testo del messaggio"]),
         uscite=["pdf_email_2_allegati.pdf"], chiamate=["smtp_sendmail"],
         controllo=lambda config: controlla_email("Due allegati", ["pdf_email_2_allegati.pdf", "allegato2.txt"]))
    caso("pdf_telegram", gui("pdf_telegram.txt", "file", "pdf_telegram.pdf", invio="telegram",
                             destinatario="telegram 123456789",
                             coda=["pdf_telegram.pdf", "Ecco il documento", "seconda riga"]),
         uscite=["pdf_telegram.pdf"], chiamate=["tg_sendMessage", "tg_sendDocument"])
    caso("pdf_whatsapp", gui("pdf_whatsapp.txt", "file", "pdf_whatsapp.pdf", invio="whatsapp",
                             destinatario="whatsapp +39 333 1234567",
                             coda=["pdf_whatsapp.pdf", "Buongiorno,", "in allegato il documento."]),
         uscite=["pdf_whatsapp.pdf"], chiamate=["wa_testo", "wa_documento"])
    caso("pdf_whatsapp_template", gui("pdf_whatsapp_template.txt", "file", "pdf_whatsapp_template.pdf",
                                      invio="whatsapp", destinatario="whatsapp 393331234567 fattura it",
                                      coda=["pdf_whatsapp_template.pdf", "Mario Rossi", "123"]),
         uscite=["pdf_whatsapp_template.pdf"], chiamate=["wa_template"])

    # --- Excel ---
    testo_excel = ("Codice;Descrizione;Quantita;Prezzo;Data;\n"
                   "1;Viti;100;0,15;20/10/2026;\n"
                   "2;Bulloni;50;1,20;21/10/2026;\n"
                   "3;Rondelle più grandi;7;12;22/10/2026;\n")
    caso("excel", gui("excel.txt", "excel", "excel.xlsx", formati=["1", "0", "1", "8", "12", "fine"]),
         txt=testo_excel, uscite=["excel.xlsx"],
         # linea 8 del .gui = colonna A, linea 9 = colonna B, ...
         controllo=controlla_excel("excel.xlsx", 4, 5, celle=[
             # .gui di prima (senza la linea "colonna"): la prima riga resta testo, senza formato
             (1, 1, "Codice", "General"), (2, 1, 1, "0"), (2, 2, "Viti", "General"), (3, 3, 50, "0"),
             (2, 4, 0.15, '#,##0.00 "€";[Red]-#,##0.00 "€"'),
             (4, 4, 12, '#,##0.00 "€";[Red]-#,##0.00 "€"'),
             (2, 5, datetime.datetime(2026, 10, 20), "dd/mm/yyyy")]))
    caso("excel_senza_fine", gui("excel_senza_fine.txt", "excel", "excel_senza_fine.xlsx",
                                 formati=["1", "0", "1", "8", "12", "0", "0"]),
         txt=testo_excel, uscite=["excel_senza_fine.xlsx"],
         controllo=controlla_excel("excel_senza_fine.xlsx", 4, 5))
    caso("excel_righe_diverse", gui("excel_righe_diverse.txt", "excel", "excel_righe_diverse.xlsx",
                                    formati=["0", "0", "1", "fine"]),
         txt="A;B;C;D;\n1;2;3;4;\n5;6;\n\n7;8;9;10;\n", uscite=["excel_righe_diverse.xlsx"],
         controllo=controlla_excel("excel_righe_diverse.xlsx", 5, 4, vuote=[(3, 3), (3, 4)]))
    caso("excel_email", gui("excel_email.txt", "excel", "excel_email.xlsx", invio="email",
                            formati=["1", "0", "1", "8", "12", "fine"],
                            coda=["cliente@example.com", "Listino", "", "In allegato il listino"]),
         txt=testo_excel, uscite=["excel_email.xlsx"], chiamate=["smtp_sendmail"])
    # Colonne con carattere, intestazione, larghezza e allineamento (linee nuove del .gui)
    # Colonne con carattere, larghezza e allineamento; righe speciali (intestazione come testo, totale
    # sull'ultima riga contata dalla fine) che valgono più della colonna
    caso("excel_formattato", gui("excel_formattato.txt", "excel", "excel_formattato.xlsx", formati=[
        "1\tnome=Intestazione\tcarattere=hsn11b\tsfondo=grigio\ttesto=si",
        "u1\tnome=Totale\tcarattere=tss14n\tsfondo=giallo\tallineamento=sinistra",
        "colonna",
        "1\tcarattere=csn12r\tlarghezza=9\tallineamento=centro",
        "0", "1", "24\tcarattere=kns10v\tallineamento=destra", "12", "fine"]),
         txt=testo_excel, uscite=["excel_formattato.xlsx"],
         controllo=controlla_excel_stili("excel_formattato.xlsx", [
             # (riga, colonna, font, dimensione, grassetto, corsivo, colore, allineamento, formato, sfondo)
             (1, 1, "Arial", 11, True, False, "FF0000FF", "center", "General", "FFD9D9D9"),
             (2, 1, "Courier New", 12, True, False, "FFFF0000", "center", "0", None),
             (3, 4, "Calibri", 10, False, True, "FF008000", "right", '#,##0.00 "€";-#,##0.00 "€"', None),
             (1, 5, "Arial", 11, True, False, "FF0000FF", None, "General", "FFD9D9D9"),
             (2, 2, "Calibri", 11, False, False, None, None, "General", None),
             # Ultima riga (u1): carattere e allineamento della riga, formato numerico della colonna
             (4, 4, "Times New Roman", 14, True, True, "FF000000", "left", '#,##0.00 "€";-#,##0.00 "€"',
              "FFFFF2CC")],
             larghezze={"A": 9}, valori={(1, 3): "Quantita", (4, 4): 12, (4, 5): datetime.datetime(2026, 10, 22)}))
    # Più intervalli di righe, contati dall'inizio e dalla fine; una riga in due stili prende l'ultimo
    caso("excel_righe_speciali", gui("excel_righe_speciali.txt", "excel", "excel_righe_speciali.xlsx", formati=[
        "1-2, u2-u1\tnome=Titoli\tcarattere=ksn11n\ttesto=si",
        "u1\tnome=Totale\tcarattere=ksn12r",
        "colonna", "0", "4", "fine"]),
         txt="Listino ottobre;;\nArticolo;Prezzo;\nViti;1.000,5;\nDadi;2,25;\nSubtotale;1002,75;\nTotale;1.002,75;\n",
         uscite=["excel_righe_speciali.xlsx"],
         controllo=controlla_excel_stili("excel_righe_speciali.xlsx", [
             (1, 1, "Calibri", 11, True, False, "FF000000", None, "General", None),
             (2, 2, "Calibri", 11, True, False, "FF000000", None, "General", None),
             (3, 2, "Calibri", 11, False, False, None, None, "#,##0.00", None),
             (5, 2, "Calibri", 11, True, False, "FF000000", None, "General", None),
             (6, 2, "Calibri", 12, True, False, "FFFF0000", None, "#,##0.00", None)],
             valori={(2, 2): "Prezzo", (3, 2): 1000.5, (5, 2): "1002,75", (6, 2): 1002.75}))

    # --- Word ---
    caso("docx", gui("docx.txt", "docx", "docx.docx", formati=FORMATI_DOCX),
         uscite=["docx.docx"],
         controllo=controlla_docx("docx.docx", TESTO_PDF))
    caso("docx_senza_formati", gui("docx_senza_formati.txt", "docx", "docx_senza_formati.docx",
                                   formati=["fine"]),
         uscite=["docx_senza_formati.docx"], controllo=controlla_docx("docx_senza_formati.docx", TESTO_PDF))
    righe_lunghe = ["Riga %d" % n for n in range(1, 13)]
    caso("docx_piu_righe_dei_formati", gui("docx_piu_righe_dei_formati.txt", "docx",
                                           "docx_piu_righe_dei_formati.docx", destinatario="telegram 1",
                                           formati=FORMATI_DOCX),
         txt="\n".join(righe_lunghe) + "\n", uscite=["docx_piu_righe_dei_formati.docx"],
         controllo=controlla_docx("docx_piu_righe_dei_formati.docx", righe_lunghe))
    caso("docx_telegram", gui("docx_telegram.txt", "docx", "docx_telegram.docx", invio="telegram",
                              destinatario="telegram 123456789", formati=FORMATI_DOCX,
                              coda=["docx_telegram.docx", "Documento Word"]),
         uscite=["docx_telegram.docx"], chiamate=["tg_sendMessage", "tg_sendDocument"])

    # --- Barcode, immagini, exe (solo .gui, niente .txt) ---
    # Percorsi completi come li scrive il gestionale, con il separatore del sistema (\ su Windows, / su Linux)
    barcode_righe = ["barcode_doc", "barcode", "2",
                     "123456789012", "ean13", "300,500", "400,600", "@LAVORO@" + os.sep + "barcode_doc.pdf", "0",
                     "ABC-123", "code128", "100,100", "300,160", "@LAVORO@" + os.sep + "barcode_doc.pdf", "0"]
    caso("barcode", barcode_righe, txt=None, lavoro_extra=["barcode_doc.pdf"],
         uscite=["@LAVORO@barcode_doc.pdf"])
    immagini_righe = ["immagini_doc", "immagini", "1",
                      "@LAVORO@" + os.sep + "logo.png", "100,100", "400,280",
                      "@LAVORO@" + os.sep + "immagini_doc.pdf", "0"]

    def controlla_immagine(config):
        doc = pymupdf.open(os.path.join(config["cartella_lavoro"], "immagini_doc.pdf"))
        box = doc[0].get_image_bbox(doc[0].get_images(full=True)[0])
        doc.close()
        # logo 200x120 nel rettangolo 300x180: deve riempirlo
        if abs(box.width - 300) > 1 or abs(box.height - 180) > 1:
            return ["immagine nel rettangolo %s invece di 100,100-400,280" % box]
        return []
    caso("immagini", immagini_righe, txt=None, lavoro_extra=["immagini_doc.pdf", "logo.png"],
         uscite=["@LAVORO@immagini_doc.pdf"], controllo=controlla_immagine)
    # Un .gui sbagliato (template inesistente) non deve bloccare quelli dopo in ordine alfabetico
    def controlla_file_dopo_errore(config):
        if not os.path.isfile(os.path.join(config["cartella_output"], "z_dopo.pdf")):
            return ["il file dopo quello sbagliato non è stato elaborato"]
        return []
    caso("errore_non_blocca", gui("errore_non_blocca.txt", "file", "a_sbagliato.pdf", "non_esiste.pdf"),
         lavoro_extra=["z_dopo.txt", "z_dopo.gui"], controllo=controlla_file_dopo_errore, rimossi=False,
         errore_atteso=True)
    # Un .bat si avvia solo su Windows
    if os.name == "nt":
        caso("exe", ["programma di prova.bat", "exe"], txt=None, lavoro_extra=["programma di prova.bat"],
             chiamate=["os_system"],
             controllo=lambda config: [] if chiamate[-1][1][0].startswith('"') else ["percorso exe senza virgolette"])

    # --- Template veri: ogni PDF della cartella dei modelli unito al testo di prova ---
    if modelli:
        for nome_modello in sorted(os.listdir(modelli)):
            if not nome_modello.lower().endswith(".pdf"):
                continue
            nome = "modello_" + os.path.splitext(nome_modello)[0]
            caso(nome, gui(nome + ".txt", "file", nome + ".pdf", nome_modello), uscite=[nome + ".pdf"],
                 controllo=testo_pdf(nome + ".pdf", "Cliente:", "città, perché, più, così"),
                 modello=os.path.join(modelli, nome_modello))
    return c


def controlla_email(oggetto, allegati, testo=None, destinatario=None):
    inviate = [d for n, d in chiamate if n == "smtp_sendmail"]
    if not inviate:
        return ["nessuna email"]
    da, a, messaggio = inviate[0]
    errori = []
    if messaggio["Subject"] != oggetto:
        errori.append("oggetto %r invece di %r" % (messaggio["Subject"], oggetto))
    if destinatario is not None and a != destinatario:
        errori.append("destinatario %r invece di %r" % (a, destinatario))
    if testo is not None:
        corpo = next(p for p in messaggio.walk() if p.get_content_type() == "text/plain")
        corpo = corpo.get_payload(decode=True).decode(corpo.get_content_charset() or "utf-8")
        if corpo != testo:
            errori.append("testo %r invece di %r" % (corpo, testo))
    nomi = [p.get_filename() for p in messaggio.walk() if p.get_filename()]
    if nomi != allegati:
        errori.append("allegati %r invece di %r" % (nomi, allegati))
    return errori


def controlla_excel(nome, righe, colonne, vuote=(), celle=()):
    """celle: (riga, colonna, valore atteso, formato atteso)"""
    def controllo(config):
        from openpyxl import load_workbook
        wb = load_workbook(os.path.join(config["cartella_output"], nome))
        ws = wb.worksheets[0]
        errori = []
        for r, c, valore, formato in celle:
            cella = ws.cell(row=r, column=c)
            if cella.value != valore or cella.number_format != formato:
                errori.append("cella %d,%d: %r %r invece di %r %r" % (r, c, cella.value, cella.number_format,
                                                                    valore, formato))
        if wb.sheetnames != ["sheet1"]:
            errori.append("fogli %r" % wb.sheetnames)
        if (ws.max_row, ws.max_column) != (righe, colonne):
            errori.append("%d righe x %d colonne invece di %d x %d" % (ws.max_row, ws.max_column, righe, colonne))
        for r, c in vuote:
            if ws.cell(row=r, column=c).value not in (None, ""):
                errori.append("cella %d,%d dovrebbe essere vuota: %r" % (r, c, ws.cell(row=r, column=c).value))
        return errori
    return controllo


def controlla_excel_stili(nome, celle, larghezze=None, valori=None):
    """celle: (riga, colonna, font, dimensione, grassetto, corsivo, colore ARGB o None, allineamento, formato,
    sfondo ARGB o None); valori: {(riga, colonna): valore atteso}"""
    def controllo(config):
        from openpyxl import load_workbook
        ws = load_workbook(os.path.join(config["cartella_output"], nome)).worksheets[0]
        errori = []
        for r, c, font, dimensione, grassetto, corsivo, colore, allineamento, formato, sfondo in celle:
            cella = ws.cell(row=r, column=c)
            # Senza colore scritto Excel usa quello del tema (rgb non è un testo)
            colore_trovato = getattr(cella.font.color, "rgb", None)
            sfondo_trovato = cella.fill.start_color.rgb if cella.fill.fill_type == "solid" else None
            trovato = (cella.font.name, cella.font.size, bool(cella.font.bold), bool(cella.font.italic),
                       colore_trovato if isinstance(colore_trovato, str) else None,
                       cella.alignment.horizontal, cella.number_format, sfondo_trovato)
            atteso = (font, dimensione, grassetto, corsivo, colore, allineamento, formato, sfondo)
            if trovato != atteso:
                errori.append("cella %d,%d: %r invece di %r" % (r, c, trovato, atteso))
        for (r, c), valore in (valori or {}).items():
            if ws.cell(row=r, column=c).value != valore:
                errori.append("cella %d,%d: valore %r invece di %r" % (r, c, ws.cell(row=r, column=c).value, valore))
        for lettera, larghezza in (larghezze or {}).items():
            if ws.column_dimensions[lettera].width != larghezza:
                errori.append("colonna %s larga %r invece di %r" % (lettera, ws.column_dimensions[lettera].width,
                                                                   larghezza))
        return errori
    return controllo


def controlla_docx(nome, righe):
    def controllo(config):
        from docx import Document
        doc = Document(os.path.join(config["cartella_output"], nome))
        testo = "\n".join(p.text for p in doc.paragraphs)
        errori = []
        for r in righe:
            if r != "" and testo.count(r + "\n") != 1:
                errori.append("%r compare %d volte nel Word" % (r, testo.count(r + "\n")))
        for p in doc.paragraphs:
            for run in p.runs:
                if run.font.name not in (None, "Times New Roman", "Arial", "Courier New"):
                    errori.append("font sconosciuto a Word: %r" % run.font.name)
                    return errori
        return errori
    return controllo


# --- Simulazioni ------------------------------------------------------------

chiamate = []


def registra(nome, *dati):
    chiamate.append((nome, dati))


class FintoSMTP:
    def __init__(self, server, porta, context=None):
        registra("smtp_connect", server, porta)

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def login(self, utente, password):
        registra("smtp_login", utente)

    def sendmail(self, da, a, testo):
        registra("smtp_sendmail", da, a, email.message_from_string(testo))


class FintoBot:
    def __init__(self, token):
        pass

    def sendMessage(self, chat_id, testo):
        registra("tg_sendMessage", chat_id, testo)

    def sendDocument(self, chat_id, document):
        registra("tg_sendDocument", chat_id, document.name)
        document.close()


class FintoWhatsApp:
    def __init__(self, token, phone_id, versione):
        pass

    def invia_testo(self, numero, testo):
        registra("wa_testo", numero, testo)

    def invia_documento(self, numero, percorso):
        registra("wa_documento", numero, percorso, os.path.isfile(percorso))

    def invia_template(self, numero, nome, lingua, parametri, documento):
        registra("wa_template", numero, nome, lingua, parametri, documento)


def carica_programma(nome_file):
    """Esegue il sorgente del programma senza l'avvio finale (start() / loop)."""
    percorso = os.path.join(CARTELLA_PROGRAMMA, nome_file + ".py")
    with open(percorso, encoding="utf-8") as f:
        sorgente = f.read()
    for fine in ("\n#Crea un array con i file precedenti", "\nstart()"):
        pos = sorgente.rfind(fine)
        if pos != -1:
            sorgente = sorgente[:pos]
            break
    sorgente += "\nlinea_inizio_format_pdf = 7\nlinea_inizio_format_excel = 7\n"
    modulo = types.ModuleType(nome_file)
    modulo.__file__ = percorso
    exec(compile(sorgente, percorso, "exec"), modulo.__dict__)
    return modulo


def prepara_simulazioni(modulo):
    modulo.smtplib = types.SimpleNamespace(SMTP_SSL=FintoSMTP)
    modulo.telepot = types.SimpleNamespace(Bot=FintoBot)
    if hasattr(modulo, "WhatsApp"):
        modulo.WhatsApp = FintoWhatsApp
    if hasattr(modulo, "stampa_pdf"):
        modulo.stampa_pdf = lambda p, *a, **k: registra("stampa_pdf", p, os.path.isfile(p))
    # Visualizzazione a schermo (os.startfile su Windows, xdg-open su Linux)
    if hasattr(modulo, "apri_file"):
        modulo.apri_file = lambda p: registra("startfile", p, os.path.isfile(p))
    finto_os = types.ModuleType("os")
    finto_os.__dict__.update(os.__dict__)
    finto_os.startfile = lambda p, *a: registra("startfile", p, os.path.isfile(p))
    finto_os.system = lambda cmd: registra("os_system", cmd)
    modulo.os = finto_os


def controlla_log(config, caso):
    """Il log deve esserci, con l'esito di ogni operazione e gli errori."""
    if "file_log" not in config:
        return []
    if not os.path.isfile(config["file_log"]):
        return ["log non creato: " + config["file_log"]]
    with open(config["file_log"], encoding="utf-8") as f:
        righe = [r for r in f.read().splitlines() if not r.startswith("    ")]
    errori = []
    for riga in righe:
        campi = riga.split(" | ")
        if len(campi) < 4:
            errori.append("riga del log non valida: " + riga)
    esiti = [r.split(" | ")[3] for r in righe if len(r.split(" | ")) >= 4]
    if caso["errore_atteso"] and "ERRORE" not in esiti:
        errori.append("l'errore non è nel log")
    if not caso["errore_atteso"]:
        if "ERRORE" in esiti:
            errori.append("errori nel log:\n" + "\n".join(r for r in righe if " | ERRORE" in r))
        elif "OK" not in esiti:
            errori.append("nessuna operazione nel log")
    return errori


def crea_template(percorso):
    doc = pymupdf.open()
    pagina = doc.new_page(width=595, height=842)
    pagina.draw_rect(pymupdf.Rect(20, 20, 575, 822), color=(0, 0, 1))
    pagina.insert_text((40, 810), "MODELLO DI PROVA", fontsize=9, color=(0, 0, 1))
    doc.save(percorso)
    doc.close()


def prepara_caso(caso):
    base = os.path.join(CARTELLA_RUN, caso["nome"])
    cartelle = {k: os.path.join(base, k) for k in ("lavoro", "template", "output")}
    for p in cartelle.values():
        os.makedirs(p)
    crea_template(os.path.join(cartelle["template"], "modello.pdf"))
    if caso["modello"]:
        shutil.copy(caso["modello"], cartelle["template"])
    for nome in caso["lavoro_extra"]:
        destinazione = os.path.join(cartelle["lavoro"], nome)
        if nome.endswith(".pdf"):
            crea_template(destinazione)
        elif nome.endswith(".png"):
            Image.new("RGB", (200, 120), (200, 30, 30)).save(destinazione)
        elif nome.endswith(".bat"):
            with open(destinazione, "w") as f:
                f.write("@echo programma di prova\n")
        elif nome.endswith(".txt"):
            with open(destinazione, "w") as f:
                f.write("Documento elaborato dopo un errore\n")
        elif nome.endswith(".gui"):
            with open(destinazione, "w") as f:
                f.write("\n".join(gui(nome[:-4] + ".txt", "file", nome[:-4] + ".pdf", formati=["fine"])) + "\n")
    allegato2 = os.path.join(base, "allegato2.txt")
    with open(allegato2, "w") as f:
        f.write("secondo allegato\n")

    righe = [r.replace("@LAVORO@", cartelle["lavoro"]).replace("@ALLEGATO2@", allegato2)
             for r in caso["gui"]]
    if caso["txt"] is not None:
        with open(os.path.join(cartelle["lavoro"], caso["nome"] + ".txt"), "w") as f:
            f.write(caso["txt"])
    with open(os.path.join(cartelle["lavoro"], caso["nome_gui"]), "w") as f:
        f.write("\n".join(righe) + "\n")
    # Copia leggibile del .gui (con i segnaposto) da riusare a mano
    with open(os.path.join(CARTELLA_CASI, caso["nome_gui"]), "w") as f:
        f.write("\n".join(caso["gui"]) + "\n")
    config = {
        "cartella_lavoro": cartelle["lavoro"], "cartella_template": cartelle["template"],
        "cartella_output": cartelle["output"], "email_mittente": "prova@example.com",
        "email_password": "x", "email_smtp_server": "smtp.example.com", "email_smtp_porta": "465",
        "telegram_token": "x", "whatsapp_token": "x", "whatsapp_phone_number_id": "1",
        "whatsapp_versione_api": "v23.0", "whatsapp_lingua": "it", "stampante": "",
        "file_log": os.path.join(base, "FileGUY.log"),
    }
    return config


def controlla_pdf(percorso):
    doc = pymupdf.open(percorso)
    try:
        testo = "".join(p.get_text() for p in doc)
        return "%d pag., %d caratteri di testo" % (doc.page_count, len(testo.strip()))
    finally:
        doc.close()


def esegui(nome_programma, modelli=None):
    shutil.rmtree(CARTELLA_RUN, ignore_errors=True)
    os.makedirs(CARTELLA_RUN)
    os.makedirs(CARTELLA_CASI, exist_ok=True)
    risultati = []
    for caso in casi(modelli):
        config = prepara_caso(caso)
        del chiamate[:]
        errori = []
        cwd = os.getcwd()
        os.chdir(CARTELLA_PROGRAMMA)
        # start() scrive su stderr gli errori dei singoli file e continua: li raccoglie
        stderr = io.StringIO()
        try:
            with contextlib.redirect_stderr(stderr):
                modulo = carica_programma(nome_programma)
                prepara_simulazioni(modulo)
                modulo.carica_config = lambda: config
                modulo.start()
        except Exception:
            errori.append(traceback.format_exc(limit=-3).strip())
        finally:
            os.chdir(cwd)
        if caso["errore_atteso"] and "Traceback" not in stderr.getvalue():
            errori.append("l'errore atteso non è stato segnalato")
        elif not caso["errore_atteso"] and "Traceback" in stderr.getvalue():
            errori.append(stderr.getvalue().strip())
        dettagli = []
        for uscita in caso["uscite"]:
            if uscita.startswith("@LAVORO@"):
                percorso = os.path.join(config["cartella_lavoro"], uscita[len("@LAVORO@"):])
            else:
                percorso = os.path.join(config["cartella_output"], uscita)
            if not os.path.isfile(percorso):
                errori.append("manca il file di uscita " + uscita)
            elif percorso.endswith(".pdf"):
                try:
                    dettagli.append(uscita + ": " + controlla_pdf(percorso))
                except Exception as e:
                    errori.append(uscita + " non è un PDF valido: " + str(e))
            else:
                dettagli.append(uscita + ": %d byte" % os.path.getsize(percorso))
        errori += controlla_log(config, caso)
        fatte = [n for n, _ in chiamate]
        for attesa in caso["chiamate"]:
            if attesa not in fatte:
                errori.append("non è stato chiamato " + attesa)
        if not errori and caso["controllo"] is not None:
            try:
                errori += caso["controllo"](config)
            except Exception:
                errori.append("controllo fallito: " + traceback.format_exc(limit=-2).strip())
        rimasti = [f for f in os.listdir(config["cartella_lavoro"])
                   if f.endswith(".gui") or f.endswith(".txt") or f == "systemgen.pdf"
                   or f.startswith("barcode") and f.endswith(".png")]
        if caso["rimossi"] and rimasti and not errori:
            errori.append("file rimasti nella cartella di lavoro: " + ", ".join(rimasti))
        risultati.append((caso["nome"], errori, dettagli, list(chiamate)))

    ok = 0
    for nome, errori, dettagli, fatte in risultati:
        stato = "OK " if not errori else "ERR"
        ok += not errori
        print("[%s] %s" % (stato, nome))
        for d in dettagli:
            print("       " + d)
        for n, dati in fatte:
            print("       -> %s %s" % (n, ", ".join(repr(x) for x in dati)[:150]))
        for e in errori:
            print("       " + e.replace("\n", "\n       "))
    print("\n%d/%d casi OK (%s)" % (ok, len(risultati), nome_programma))
    return ok == len(risultati)


if __name__ == "__main__":
    sys.exit(0 if esegui(sys.argv[1] if len(sys.argv) > 1 else "SingleGUY",
                         sys.argv[2] if len(sys.argv) > 2 else None) else 1)
